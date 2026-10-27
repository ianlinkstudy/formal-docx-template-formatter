from __future__ import annotations

import re
import shutil
import uuid
from collections import Counter
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .common import (
    ANCHOR_RE,
    DEFAULT_LIBRARY_DIR,
    GB_PROFILE_PATH,
    MAX_CUSTOM_TEMPLATES,
    ROLE_NAMES,
    FormatterError,
    StateError,
    atomic_copy,
    atomic_write_json,
    deep_merge,
    extract_anchor_names,
    hash_file,
    library_lock,
    load_json,
    normalize_anchor,
    open_document,
    prepare_docx_copy,
    role_map,
    section_values,
    stable_hash,
    utc_now,
)


def load_gb_profile() -> dict[str, Any]:
    return load_json(GB_PROFILE_PATH)


def _ensure_layout(library_dir: Path) -> None:
    for name in ("pending", "templates", "archive"):
        (library_dir / name).mkdir(parents=True, exist_ok=True)


def _registry_path(library_dir: Path) -> Path:
    return library_dir / "registry.json"


def _initial_registry() -> dict[str, Any]:
    return {"schema_version": 1, "active": [], "archived": []}


def load_registry(library_dir: Path = DEFAULT_LIBRARY_DIR) -> dict[str, Any]:
    _ensure_layout(library_dir)
    path = _registry_path(library_dir)
    if not path.exists():
        return _initial_registry()
    registry = load_json(path)
    registry.setdefault("active", [])
    registry.setdefault("archived", [])
    if not isinstance(registry["active"], list) or not isinstance(
        registry["archived"], list
    ):
        raise FormatterError("模板注册表格式损坏。")
    return registry


def save_registry(
    registry: dict[str, Any], library_dir: Path = DEFAULT_LIBRARY_DIR
) -> None:
    atomic_write_json(_registry_path(library_dir), registry)


def name_key(name: str) -> str:
    return name.strip().casefold()


def validate_template_name(name: str) -> str:
    cleaned = name.strip()
    if not cleaned:
        raise StateError("模板名称不能为空。")
    if len(cleaned) > 32:
        raise StateError("模板名称最多 32 个字符。")
    if any(character in cleaned for character in ("/", "\\", "\0")):
        raise StateError("模板名称不能包含路径分隔符。")
    if any(ord(character) < 32 for character in cleaned):
        raise StateError("模板名称不能包含控制字符。")
    if cleaned.casefold() in {"gb9704", "gb", "gb/t 9704-2012 参考版"}:
        raise StateError("该名称属于内置模板，请换一个名称。")
    return cleaned


def validate_profile(profile: dict[str, Any]) -> None:
    if profile.get("schema_version") != 2:
        raise StateError("模板配置版本必须为 2。")
    page = profile.get("page")
    if not isinstance(page, dict):
        raise StateError("模板配置缺少 page。")
    for key in ("width_mm", "height_mm"):
        if not isinstance(page.get(key), (int, float)) or page[key] <= 0:
            raise StateError(f"page.{key} 必须是正数。")
    for key in ("margin_top_mm", "margin_bottom_mm", "margin_left_mm", "margin_right_mm"):
        if not isinstance(page.get(key), (int, float)) or page[key] < 0:
            raise StateError(f"page.{key} 必须是非负数。")
    styles = profile.get("styles")
    if not isinstance(styles, dict):
        raise StateError("模板配置缺少 styles。")
    for role in ROLE_NAMES:
        style = styles.get(role)
        if not isinstance(style, dict):
            raise StateError(f"缺少样式角色：{role}")
        for key in ("font", "latin_font"):
            if not isinstance(style.get(key), str) or not style[key].strip():
                raise StateError(f"styles.{role}.{key} 不能为空。")
        size = style.get("size_pt")
        if not isinstance(size, (int, float)) or not 5 <= float(size) <= 72:
            raise StateError(f"styles.{role}.size_pt 必须在 5–72 之间。")
        if style.get("bold") is not None and not isinstance(style["bold"], bool):
            raise StateError(f"styles.{role}.bold 必须是布尔值或 null。")
        if style.get("italic") is not None and not isinstance(
            style["italic"], bool
        ):
            raise StateError(f"styles.{role}.italic 必须是布尔值或 null。")
        if style.get("alignment") not in {
            "left",
            "center",
            "right",
            "justify",
            "distribute",
        }:
            raise StateError(f"styles.{role}.alignment 无效。")
        if not re.fullmatch(r"[0-9A-Fa-f]{6}", str(style.get("color", ""))):
            raise StateError(f"styles.{role}.color 必须是六位颜色代码。")
        line_rule = style.get("line_rule")
        if line_rule not in {"exact", "at_least", "multiple", "single"}:
            raise StateError(f"styles.{role}.line_rule 无效。")
        if line_rule != "single":
            value = style.get("line_value")
            if not isinstance(value, (int, float)) or value <= 0:
                raise StateError(f"styles.{role}.line_value 必须是正数。")
        for key in ("space_before_pt", "space_after_pt"):
            value = style.get(key)
            if not isinstance(value, (int, float)) or value < 0:
                raise StateError(f"styles.{role}.{key} 必须是非负数。")
        for key in ("first_line_chars", "right_chars"):
            value = style.get(key)
            if not isinstance(value, (int, float)) or value < 0:
                raise StateError(f"styles.{role}.{key} 必须是非负数。")
    table = profile.get("table")
    if not isinstance(table, dict):
        raise StateError("模板配置缺少 table。")
    for key in ("header_size_pt", "body_size_pt"):
        if not isinstance(table.get(key), (int, float)) or not 5 <= float(
            table[key]
        ) <= 72:
            raise StateError(f"table.{key} 必须在 5–72 之间。")
    if not isinstance(table.get("border_size_eighths_pt"), int) or not (
        1 <= table["border_size_eighths_pt"] <= 48
    ):
        raise StateError("table.border_size_eighths_pt 必须在 1–48 之间。")
    if not re.fullmatch(r"[0-9A-Fa-f]{6}", str(table.get("border_color", ""))):
        raise StateError("table.border_color 必须是六位颜色代码。")
    for edge in ("top", "bottom", "left", "right"):
        key = f"cell_margin_{edge}_twips"
        if not isinstance(table.get(key), int) or table[key] < 0:
            raise StateError(f"table.{key} 必须是非负整数。")
    figure = profile.get("figure")
    if not isinstance(figure, dict):
        raise StateError("模板配置缺少 figure。")
    ratio = figure.get("max_width_ratio")
    if not isinstance(ratio, (int, float)) or not 0.1 <= float(ratio) <= 1.0:
        raise StateError("figure.max_width_ratio 必须在 0.1–1.0 之间。")


def template_choices(library_dir: Path = DEFAULT_LIBRARY_DIR) -> dict[str, Any]:
    registry = load_registry(library_dir)
    active = sorted(registry["active"], key=lambda item: item["created_at"])
    return {
        "builtin": [
            {
                "id": "gb9704",
                "name": "GB/T 9704-2012 参考版",
                "mode": "style",
                "description": "内置中文正式文档参考版式，不占自定义模板槽位。",
            }
        ],
        "custom": [
            {
                "id": item["id"],
                "name": item["name"],
                "mode": item["mode"],
                "created_at": item["created_at"],
                "confirmed_at": item["confirmed_at"],
            }
            for item in active
        ],
        "active_custom_count": len(active),
        "remaining_slots": MAX_CUSTOM_TEMPLATES - len(active),
        "max_custom_templates": MAX_CUSTOM_TEMPLATES,
        "can_import": len(active) < MAX_CUSTOM_TEMPLATES,
        "requires_selection": True,
    }


def _candidate_dir(candidate_id: str, library_dir: Path) -> Path:
    if not re.fullmatch(r"[a-f0-9]{32}", candidate_id):
        raise StateError("候选 ID 无效。")
    path = library_dir / "pending" / candidate_id
    if not path.is_dir():
        raise StateError("候选不存在或已过期。")
    return path


def load_candidate(
    candidate_id: str, library_dir: Path = DEFAULT_LIBRARY_DIR
) -> tuple[Path, dict[str, Any]]:
    directory = _candidate_dir(candidate_id, library_dir)
    return directory, load_json(directory / "candidate.json")


def _save_candidate(directory: Path, candidate: dict[str, Any]) -> None:
    candidate["candidate_sha256"] = _candidate_hash(candidate)
    atomic_write_json(directory / "candidate.json", candidate)


def _candidate_hash(candidate: dict[str, Any]) -> str:
    payload = {
        "mode": candidate["mode"],
        "profile": candidate["profile"],
        "shell": candidate["shell"],
        "source_sha256": candidate["source_sha256"],
    }
    return stable_hash(payload)


def _modal(values: list[Any], default: Any = None) -> tuple[Any, float]:
    if not values:
        return default, 0.0
    encoded = [stable_hash(value) for value in values]
    counts = Counter(encoded)
    winner_hash, winner_count = counts.most_common(1)[0]
    winner = values[encoded.index(winner_hash)]
    return winner, winner_count / len(values)


def _run_east_asia(run: Any) -> str | None:
    r_pr = run._element.rPr
    if r_pr is not None and r_pr.rFonts is not None:
        return r_pr.rFonts.get(qn("w:eastAsia"))
    return None


def _run_latin(run: Any) -> str | None:
    r_pr = run._element.rPr
    if r_pr is not None and r_pr.rFonts is not None:
        return r_pr.rFonts.get(qn("w:ascii"))
    return run.font.name


def _style_font(paragraph: Paragraph, east_asia: bool = True) -> str | None:
    style = paragraph.style
    if style is None:
        return None
    if not east_asia:
        return style.font.name
    style_element = style.element
    r_pr = style_element.find(qn("w:rPr"))
    if r_pr is None:
        return style.font.name
    fonts = r_pr.find(qn("w:rFonts"))
    if fonts is None:
        return style.font.name
    return fonts.get(qn("w:eastAsia")) or style.font.name


def _effective_run_style(paragraph: Paragraph, base: dict[str, Any]) -> dict[str, Any]:
    text_runs = [run for run in paragraph.runs if run.text.strip()]
    runs = text_runs or list(paragraph.runs)
    fonts: list[str] = []
    latin_fonts: list[str] = []
    sizes: list[float] = []
    bolds: list[bool | None] = []
    italics: list[bool | None] = []
    colors: list[str] = []
    for run in runs:
        east = _run_east_asia(run) or _style_font(paragraph, True) or base["font"]
        latin = (
            _run_latin(run) or _style_font(paragraph, False) or base["latin_font"]
        )
        size = run.font.size or (
            paragraph.style.font.size if paragraph.style is not None else None
        )
        size_pt = size.pt if size is not None else float(base["size_pt"])
        bold = run.bold
        if bold is None and paragraph.style is not None:
            bold = paragraph.style.font.bold
        italic = run.italic
        if italic is None and paragraph.style is not None:
            italic = paragraph.style.font.italic
        color = run.font.color
        color_value = base.get("color", "000000")
        if color is not None and color.rgb is not None:
            color_value = str(color.rgb)
        fonts.append(east)
        latin_fonts.append(latin)
        sizes.append(round(float(size_pt), 2))
        bolds.append(bold)
        italics.append(italic)
        colors.append(color_value.upper())
    font, font_conf = _modal(fonts, base["font"])
    latin, latin_conf = _modal(latin_fonts, base["latin_font"])
    size, size_conf = _modal(sizes, base["size_pt"])
    bold, _ = _modal(bolds, base.get("bold"))
    italic, _ = _modal(italics, base.get("italic"))
    color, _ = _modal(colors, base.get("color", "000000"))
    return {
        "font": font,
        "latin_font": latin,
        "size_pt": size,
        "bold": bold,
        "italic": italic,
        "color": color,
        "_confidence": min(font_conf, latin_conf, size_conf) if runs else 0.0,
    }


def _alignment_name(value: Any, fallback: str) -> str:
    return {
        WD_ALIGN_PARAGRAPH.LEFT: "left",
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
        WD_ALIGN_PARAGRAPH.DISTRIBUTE: "distribute",
    }.get(value, fallback)


def _first_line_chars(paragraph: Paragraph, size_pt: float) -> float:
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return 0.0
    ind = p_pr.find(qn("w:ind"))
    if ind is None:
        return 0.0
    chars = ind.get(qn("w:firstLineChars"))
    if chars is not None:
        return round(int(chars) / 100, 2)
    value = ind.get(qn("w:firstLine"))
    if value is None:
        return 0.0
    return round(int(value) / max(1.0, size_pt * 20), 2)


def _right_chars(paragraph: Paragraph, size_pt: float) -> float:
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return 0.0
    ind = p_pr.find(qn("w:ind"))
    if ind is None:
        return 0.0
    chars = ind.get(qn("w:rightChars"))
    if chars is not None:
        return round(int(chars) / 100, 2)
    value = ind.get(qn("w:right"))
    if value is None:
        return 0.0
    return round(int(value) / max(1.0, size_pt * 20), 2)


def _effective_paragraph_style(
    paragraph: Paragraph, base: dict[str, Any]
) -> tuple[dict[str, Any], float]:
    run_style = _effective_run_style(paragraph, base)
    size_pt = float(run_style["size_pt"])
    fmt = paragraph.paragraph_format
    style_fmt = paragraph.style.paragraph_format if paragraph.style is not None else None
    alignment = fmt.alignment
    if alignment is None and style_fmt is not None:
        alignment = style_fmt.alignment
    rule = fmt.line_spacing_rule
    spacing = fmt.line_spacing
    if rule is None and style_fmt is not None:
        rule = style_fmt.line_spacing_rule
    if spacing is None and style_fmt is not None:
        spacing = style_fmt.line_spacing
    if isinstance(spacing, (float, int)):
        line_rule = "multiple"
        line_value = round(float(spacing), 2)
    else:
        line_rule = (
            "at_least"
            if rule == WD_LINE_SPACING.AT_LEAST
            else "exact"
            if rule == WD_LINE_SPACING.EXACTLY
            else "single"
            if rule in {WD_LINE_SPACING.SINGLE, WD_LINE_SPACING.ONE_POINT_FIVE}
            else base.get("line_rule", "exact")
        )
        line_value = (
            round(spacing.pt, 2)
            if hasattr(spacing, "pt")
            else float(base.get("line_value", 28))
        )
    before = fmt.space_before if fmt.space_before is not None else (
        style_fmt.space_before if style_fmt is not None else None
    )
    after = fmt.space_after if fmt.space_after is not None else (
        style_fmt.space_after if style_fmt is not None else None
    )
    keep = fmt.keep_with_next
    confidence = run_style.pop("_confidence")
    result = {
        **base,
        **run_style,
        "alignment": _alignment_name(alignment, base.get("alignment", "justify")),
        "line_rule": line_rule,
        "line_value": line_value,
        "space_before_pt": round(before.pt, 2) if before is not None else 0.0,
        "space_after_pt": round(after.pt, 2) if after is not None else 0.0,
        "first_line_chars": _first_line_chars(paragraph, size_pt),
        "right_chars": _right_chars(paragraph, size_pt),
        "keep_with_next": bool(keep) if keep is not None else False,
    }
    return result, confidence


def _extract_page(doc: DocumentObject) -> tuple[dict[str, Any], list[str]]:
    warnings: list[str] = []
    sections = [section_values(section) for section in doc.sections]
    if not sections:
        raise FormatterError("模板没有可读取的页面节。")
    first = sections[0]
    if len(sections) > 1:
        warnings.append(f"模板包含 {len(sections)} 个节；样式模式只使用第一节页面参数。")
    return {
        "section_strategy": "uniform",
        "width_mm": first["width_mm"],
        "height_mm": first["height_mm"],
        "margin_top_mm": first["margin_top_mm"],
        "margin_bottom_mm": first["margin_bottom_mm"],
        "margin_left_mm": first["margin_left_mm"],
        "margin_right_mm": first["margin_right_mm"],
        "source_sections": sections,
    }, warnings


def _extract_table_style(
    doc: DocumentObject,
    table_base: dict[str, Any],
    body_style: dict[str, Any],
) -> dict[str, Any]:
    result = deepcopy(table_base)
    if not doc.tables:
        return result
    table = doc.tables[0]
    if not table.rows:
        return result
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders")) if tbl_pr is not None else None
    top_border = borders.find(qn("w:top")) if borders is not None else None
    if top_border is not None:
        size = top_border.get(qn("w:sz"))
        color = top_border.get(qn("w:color"))
        if size and size.isdigit():
            result["border_size_eighths_pt"] = int(size)
        if color:
            result["border_color"] = color.upper()
    cell_margins = tbl_pr.find(qn("w:tblCellMar")) if tbl_pr is not None else None
    if cell_margins is not None:
        for edge in ("top", "left", "bottom", "right"):
            margin = cell_margins.find(qn(f"w:{edge}"))
            if margin is not None and (margin.get(qn("w:w")) or "").isdigit():
                result[f"cell_margin_{edge}_twips"] = int(margin.get(qn("w:w")))
    first_cell = table.rows[0].cells[0]
    shading = first_cell._tc.tcPr.find(qn("w:shd")) if first_cell._tc.tcPr is not None else None
    if shading is not None:
        result["header_fill"] = shading.get(qn("w:fill"))
    tr_pr = table.rows[0]._tr.trPr
    result["repeat_header_row"] = bool(
        tr_pr is not None and tr_pr.find(qn("w:tblHeader")) is not None
    )
    for row_index, row in enumerate(table.rows[: max(1, min(2, len(table.rows)))]):
        sizes: list[float] = []
        fonts: list[str] = []
        bold: list[bool | None] = []
        alignments: list[str] = []
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                style, _ = _effective_paragraph_style(
                    paragraph,
                    {
                        **body_style,
                        "font": table_base["body_font"],
                        "size_pt": table_base["body_size_pt"],
                    },
                )
                fonts.append(style["font"])
                sizes.append(float(style["size_pt"]))
                bold.append(style["bold"])
                alignments.append(style["alignment"])
        font, _ = _modal(fonts, result["body_font"])
        size, _ = _modal(sizes, result["body_size_pt"])
        weight, _ = _modal(bold, result["body_bold"])
        if row_index == 0:
            result["header_font"] = font
            result["header_size_pt"] = size
            result["header_bold"] = weight
        else:
            result["body_font"] = font
            result["body_size_pt"] = size
            result["body_bold"] = weight
    alignment, _ = _modal(alignments, result["alignment"])
    result["alignment"] = alignment
    return result


def extract_profile_from_document(doc: DocumentObject) -> tuple[dict[str, Any], dict[str, Any]]:
    base = load_gb_profile()
    page, page_warnings = _extract_page(doc)
    mapping = role_map(doc)
    styles: dict[str, Any] = {}
    role_confidence: dict[str, float] = {}
    warnings = list(page_warnings)
    for role in ROLE_NAMES:
        default_rule = deepcopy(base["styles"].get(role, base["styles"]["body"]))
        indices = mapping.get(role, [])
        if not indices:
            styles[role] = default_rule
            role_confidence[role] = 0.0
            continue
        candidates: list[dict[str, Any]] = []
        confidences: list[float] = []
        for index in indices[:24]:
            candidate, confidence = _effective_paragraph_style(
                doc.paragraphs[index], default_rule
            )
            candidates.append(candidate)
            confidences.append(confidence)
        winner, agreement = _modal(candidates, default_rule)
        styles[role] = winner
        role_confidence[role] = min(
            agreement,
            sum(confidences) / len(confidences) if confidences else 0.0,
        )
        if agreement < 0.7:
            warnings.append(f"{role} 存在多种样式，只采用出现最多的一组。")
    table = _extract_table_style(doc, base["table"], base["styles"]["body"])
    profile = {
        "schema_version": 2,
        "id": "custom",
        "name": "待确认自定义模板",
        "description": "从用户 DOCX 提取，尚未确认。",
        "mode": "style",
        "provenance": {
            "kind": "user_docx",
            "confidence": "low"
            if min(role_confidence.values(), default=0.0) < 0.5
            else "medium",
        },
        "page": page,
        "styles": styles,
        "table": table,
        "figure": deepcopy(base["figure"]),
        "formula": deepcopy(base["formula"]),
    }
    extraction = {
        "role_samples": {role: len(indices) for role, indices in mapping.items()},
        "role_confidence": {
            role: round(value, 3) for role, value in role_confidence.items()
        },
        "warnings": warnings,
    }
    return profile, extraction


def _anchor_state(doc: DocumentObject) -> tuple[dict[str, list[int]], list[str], list[str]]:
    anchors: dict[str, list[int]] = {}
    unknown: list[str] = []
    warnings: list[str] = []
    for index, paragraph in enumerate(doc.paragraphs):
        for raw in extract_anchor_names([paragraph]):
            canonical = normalize_anchor(raw)
            if canonical is None:
                unknown.append(raw)
                continue
            anchors.setdefault(canonical, []).append(index)
    top_level_count = sum(len(indices) for indices in anchors.values()) + len(unknown)
    all_xml = [doc.element.body.xml]
    for section in doc.sections:
        all_xml.extend(
            [
                section.header._element.xml,
                section.first_page_header._element.xml,
                section.even_page_header._element.xml,
                section.footer._element.xml,
                section.first_page_footer._element.xml,
                section.even_page_footer._element.xml,
            ]
        )
    all_anchor_count = sum(
        len(ANCHOR_RE.findall(xml)) for xml in all_xml
    )
    if all_anchor_count > top_level_count:
        unknown.append("锚点只能放在正文顶层普通段落，不能放在表格、页眉或页脚")
    for canonical, indices in anchors.items():
        if len(indices) > 1:
            warnings.append(f"锚点 {canonical} 出现了 {len(indices)} 次。")
    return anchors, unknown, warnings


def _template_object_summary(doc: DocumentObject) -> dict[str, Any]:
    xml = doc.element.xml
    return {
        "sections": len(doc.sections),
        "tables": len(doc.tables),
        "inline_images": len(doc.inline_shapes),
        "formula_count": xml.count("<m:oMath"),
        "object_count": xml.count("<w:object"),
        "has_header_text": any(
            p.text.strip() for section in doc.sections for p in section.header.paragraphs
        ),
        "has_footer_text": any(
            p.text.strip() for section in doc.sections for p in section.footer.paragraphs
        ),
    }


def _block_boundaries(doc: DocumentObject) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for index, block in enumerate(
        child
        for child in doc.element.body.iterchildren()
        if child.tag in {qn("w:p"), qn("w:tbl")}
    ):
        if block.tag == qn("w:p"):
            paragraph = Paragraph(block, doc)
            result.append(
                {"index": index, "kind": "paragraph", "text": paragraph.text[:120]}
            )
        else:
            table = Table(block, doc)
            sample = " ".join(
                cell.text.replace("\n", " ")
                for row in table.rows[:1]
                for cell in row.cells
            )[:120]
            result.append({"index": index, "kind": "table", "text": sample})
    return result


def _shell_summary(
    doc: DocumentObject,
    mode: str,
    body_start: int | None,
    body_end: int | None,
) -> tuple[dict[str, Any], list[str], list[str]]:
    anchors, unknown, warnings = _anchor_state(doc)
    blockers: list[str] = []
    if unknown:
        blockers.append("存在未知锚点：" + "、".join(sorted(set(unknown))))
    if len(anchors.get("BODY", [])) > 1:
        blockers.append("{{BODY}} 只能出现一次。")
    for canonical in ("TITLE", "RECIPIENT", "SIGNATURE", "DATE"):
        if len(anchors.get(canonical, [])) > 1:
            blockers.append(f"{{{{{canonical}}}}} 只能出现一次。")
    boundaries = _block_boundaries(doc)
    if mode == "shell":
        if not anchors.get("BODY") and body_start is None:
            blockers.append(
                "母版没有 {{BODY}} 锚点；请提供正文起始块，或改用样式模式。"
            )
        if body_start is not None and not (0 <= body_start < len(boundaries)):
            blockers.append("正文起始块索引超出范围。")
        if body_end is not None and body_start is None:
            blockers.append("指定正文结束块时必须同时指定起始块。")
        if (
            body_start is not None
            and body_end is not None
            and not (body_start <= body_end < len(boundaries))
        ):
            blockers.append("正文替换范围无效：结束块必须不小于起始块。")
    return {
        "anchors": anchors,
        "body_start": body_start,
        "body_end": body_end,
        "blocks": boundaries,
        "objects": _template_object_summary(doc),
    }, warnings, blockers


def _build_summary(
    candidate: dict[str, Any], extraction: dict[str, Any]
) -> dict[str, Any]:
    profile = candidate["profile"]
    low_confidence = [
        role
        for role, value in extraction.get("role_confidence", {}).items()
        if value < 0.5
    ]
    shell = candidate["shell"]
    return {
        "mode": candidate["mode"],
        "mode_label": "样式学习" if candidate["mode"] == "style" else "母版套用",
        "page": {
            key: profile["page"][key]
            for key in (
                "width_mm",
                "height_mm",
                "margin_top_mm",
                "margin_bottom_mm",
                "margin_left_mm",
                "margin_right_mm",
            )
        },
        "roles": {
            role: {
                key: profile["styles"][role][key]
                for key in (
                    "font",
                    "size_pt",
                    "bold",
                    "alignment",
                    "line_rule",
                    "line_value",
                    "first_line_chars",
                )
            }
            for role in ROLE_NAMES
        },
        "table": profile["table"],
        "objects": shell.get("objects", {}),
        "anchors": shell.get("anchors", {}),
        "body_range": {
            "start": shell.get("body_start"),
            "end": shell.get("body_end"),
        },
        "role_samples": extraction.get("role_samples", {}),
        "role_confidence": extraction.get("role_confidence", {}),
        "low_confidence_roles": low_confidence,
    }


def inspect_template(
    source: Path,
    mode: str = "auto",
    body_start: int | None = None,
    body_end: int | None = None,
    library_dir: Path = DEFAULT_LIBRARY_DIR,
) -> dict[str, Any]:
    if mode not in {"auto", "style", "shell"}:
        raise StateError("mode 必须是 auto、style 或 shell。")
    _ensure_layout(library_dir)
    choices = template_choices(library_dir)
    if choices["remaining_slots"] <= 0:
        names = "、".join(item["name"] for item in choices["custom"])
        raise StateError(f"自定义模板已满三套（{names}）。请先归档一套再导入。")

    candidate_id = uuid.uuid4().hex
    directory = library_dir / "pending" / candidate_id
    directory.mkdir(parents=True, exist_ok=False)
    try:
        working = prepare_docx_copy(source, directory / "source.docx")
        doc = open_document(working)
        anchors, _, _ = _anchor_state(doc)
        if mode == "auto":
            resolved_mode = "shell" if anchors or body_start is not None else "style"
        else:
            resolved_mode = mode
        profile, extraction = extract_profile_from_document(doc)
        validate_profile(profile)
        profile["mode"] = resolved_mode
        shell, shell_warnings, blockers = _shell_summary(
            doc, resolved_mode, body_start, body_end
        )
        if resolved_mode == "style":
            shell["anchors"] = {}
            shell["body_start"] = None
            shell["body_end"] = None
            shell["blocks"] = []
            blockers = []
        candidate = {
            "schema_version": 1,
            "candidate_id": candidate_id,
            "status": "parsed",
            "source_name": source.name,
            "source_sha256": hash_file(source),
            "working_sha256": hash_file(working),
            "created_at": utc_now(),
            "mode": resolved_mode,
            "profile": profile,
            "shell": shell,
            "extraction": extraction,
            "warnings": extraction["warnings"] + shell_warnings,
            "blockers": blockers,
            "confirmation": None,
            "candidate_sha256": "",
        }
        candidate["summary"] = _build_summary(candidate, extraction)
        _save_candidate(directory, candidate)
        return {
            "status": "needs_confirmation",
            "candidate": candidate,
        }
    except Exception:
        if directory.exists():
            shutil.rmtree(directory)
        raise


def _refresh_shell(
    directory: Path,
    candidate: dict[str, Any],
    body_start: int | None,
    body_end: int | None,
) -> None:
    doc = open_document(directory / "source.docx")
    shell, warnings, blockers = _shell_summary(
        doc, candidate["mode"], body_start, body_end
    )
    if candidate["mode"] == "style":
        shell = {
            **shell,
            "anchors": {},
            "body_start": None,
            "body_end": None,
            "blocks": [],
        }
        blockers = []
    candidate["shell"] = shell
    candidate["warnings"] = (
        candidate["extraction"]["warnings"]
        + [item for item in warnings if item not in candidate["extraction"]["warnings"]]
    )
    candidate["blockers"] = blockers


def amend_candidate(
    candidate_id: str,
    changes: dict[str, Any] | None = None,
    mode: str | None = None,
    body_start: int | None = None,
    body_end: int | None = None,
    library_dir: Path = DEFAULT_LIBRARY_DIR,
) -> dict[str, Any]:
    directory, candidate = load_candidate(candidate_id, library_dir)
    changes = deepcopy(changes or {})
    if mode is not None:
        candidate["mode"] = mode
        candidate["profile"]["mode"] = mode
    profile_changes = changes.pop("profile", None)
    if profile_changes:
        candidate["profile"] = deep_merge(candidate["profile"], profile_changes)
    shell_changes = changes.pop("shell", None)
    if shell_changes:
        candidate["shell"] = deep_merge(candidate["shell"], shell_changes)
        body_start = (
            candidate["shell"].get("body_start")
            if body_start is None
            else body_start
        )
        body_end = (
            candidate["shell"].get("body_end") if body_end is None else body_end
        )
    if changes:
        for key in ("warnings", "blockers"):
            if key in changes:
                candidate[key] = changes[key]
    candidate["status"] = "parsed"
    candidate["confirmation"] = None
    _refresh_shell(directory, candidate, body_start, body_end)
    candidate["profile"]["id"] = "custom"
    candidate["profile"]["name"] = "待确认自定义模板"
    validate_profile(candidate["profile"])
    candidate["summary"] = _build_summary(candidate, candidate["extraction"])
    _save_candidate(directory, candidate)
    return {"status": "needs_confirmation", "candidate": candidate}


def confirm_candidate(
    candidate_id: str,
    candidate_sha256: str,
    confirmation_text: str,
    library_dir: Path = DEFAULT_LIBRARY_DIR,
) -> dict[str, Any]:
    directory, candidate = load_candidate(candidate_id, library_dir)
    current = _candidate_hash(candidate)
    if candidate_sha256 != current:
        raise StateError("候选内容已变化或哈希不匹配；请重新查看摘要并确认。")
    if candidate["blockers"]:
        raise StateError("候选仍有阻断项，不能确认：" + "；".join(candidate["blockers"]))
    validate_profile(candidate["profile"])
    cleaned_confirmation = confirmation_text.strip()
    if not re.search(r"确认|同意|confirm|approve", cleaned_confirmation, re.IGNORECASE):
        raise StateError("确认文本必须明确表达“确认”或“同意”。")
    candidate["status"] = "confirmed"
    candidate["confirmation"] = {
        "confirmed_at": utc_now(),
        "candidate_sha256": current,
        "explicit": True,
        "user_text": cleaned_confirmation,
    }
    candidate["candidate_sha256"] = current
    atomic_write_json(directory / "candidate.json", candidate)
    return {"status": "confirmed", "candidate_id": candidate_id}


def _block_elements(doc: DocumentObject) -> list[Any]:
    return [
        child
        for child in doc.element.body.iterchildren()
        if child.tag in {qn("w:p"), qn("w:tbl")}
    ]


def _materialize_shell(source: Path, target: Path, shell: dict[str, Any]) -> None:
    atomic_copy(source, target)
    if shell.get("body_start") is None or shell.get("anchors", {}).get("BODY"):
        return
    doc = open_document(target)
    blocks = _block_elements(doc)
    start = int(shell["body_start"])
    end = int(shell.get("body_end") or start)
    if not (0 <= start <= end < len(blocks)):
        raise StateError("母版正文替换范围无效，未保存模板。")
    anchor = doc.add_paragraph("{{BODY}}")
    blocks[start].addprevious(anchor._p)
    for block in blocks[start : end + 1]:
        if block.getparent() is not None:
            block.getparent().remove(block)
    doc.save(target)


def _copy_template_assets(
    candidate_dir: Path,
    target_dir: Path,
    candidate: dict[str, Any],
    name: str,
) -> dict[str, Any]:
    target_dir.mkdir(parents=True, exist_ok=False)
    profile = deepcopy(candidate["profile"])
    profile["id"] = target_dir.name
    profile["name"] = name
    profile["description"] = f"用户确认并命名的本地模板：{name}"
    profile["provenance"] = {
        **profile.get("provenance", {}),
        "source_name": candidate["source_name"],
        "source_sha256": candidate["source_sha256"],
        "confirmed_at": candidate["confirmation"]["confirmed_at"],
    }
    atomic_write_json(target_dir / "profile.json", profile)
    if candidate["mode"] == "shell":
        _materialize_shell(
            candidate_dir / "source.docx", target_dir / "shell.docx", candidate["shell"]
        )
    metadata = {
        "schema_version": 1,
        "id": target_dir.name,
        "name": name,
        "name_key": name_key(name),
        "mode": candidate["mode"],
        "created_at": utc_now(),
        "confirmed_at": candidate["confirmation"]["confirmed_at"],
        "source_sha256": candidate["source_sha256"],
        "candidate_sha256": candidate["candidate_sha256"],
        "profile_file": "profile.json",
        "shell_file": "shell.docx" if candidate["mode"] == "shell" else None,
        "summary": candidate["summary"],
        "warnings": candidate["warnings"],
    }
    atomic_write_json(target_dir / "metadata.json", metadata)
    return metadata


def save_candidate(
    candidate_id: str,
    name: str,
    library_dir: Path = DEFAULT_LIBRARY_DIR,
) -> dict[str, Any]:
    directory, candidate = load_candidate(candidate_id, library_dir)
    if not candidate.get("confirmation"):
        raise StateError("模板尚未确认，不能命名或保存。")
    current_hash = _candidate_hash(candidate)
    if candidate["confirmation"]["candidate_sha256"] != current_hash:
        raise StateError("模板在确认后又发生变化；请重新确认。")
    cleaned_name = validate_template_name(name)
    validate_profile(candidate["profile"])
    with library_lock(library_dir):
        registry = load_registry(library_dir)
        if len(registry["active"]) >= MAX_CUSTOM_TEMPLATES:
            raise StateError("自定义模板已达三套，请先归档一套。")
        all_names = {item["name_key"] for item in registry["active"]}
        if name_key(cleaned_name) in all_names:
            raise StateError("模板名称已存在。")
        template_id = uuid.uuid4().hex
        target_dir = library_dir / "templates" / template_id
        try:
            metadata = _copy_template_assets(
                directory, target_dir, candidate, cleaned_name
            )
            registry["active"].append(metadata)
            save_registry(registry, library_dir)
        except Exception:
            if target_dir.exists():
                shutil.rmtree(target_dir)
            raise
    shutil.rmtree(directory)
    return {
        "status": "saved",
        "template": metadata,
        "remaining_slots": MAX_CUSTOM_TEMPLATES - len(registry["active"]),
        "next_step": "模板已保存。请在下一次排版任务开始时选择 GB 或该模板。",
    }


def find_registry_entry(
    name: str, library_dir: Path = DEFAULT_LIBRARY_DIR
) -> dict[str, Any] | None:
    registry = load_registry(library_dir)
    wanted = name_key(name)
    return next(
        (item for item in registry["active"] if item["name_key"] == wanted), None
    )


def get_template(
    name: str, library_dir: Path = DEFAULT_LIBRARY_DIR
) -> dict[str, Any]:
    if name_key(name) in {"gb9704", "gb", "gb/t 9704-2012 参考版"}:
        return {
            "id": "gb9704",
            "name": "GB/T 9704-2012 参考版",
            "mode": "style",
            "profile": load_gb_profile(),
            "profile_path": GB_PROFILE_PATH,
            "shell_path": None,
            "builtin": True,
        }
    entry = find_registry_entry(name, library_dir)
    if entry is None:
        raise StateError("未找到已保存模板。请先使用 template_cli.py list 选择。")
    base = library_dir / "templates" / entry["id"]
    result = {
        **entry,
        "profile": load_json(base / entry["profile_file"]),
        "profile_path": base / entry["profile_file"],
        "shell_path": base / entry["shell_file"] if entry.get("shell_file") else None,
        "builtin": False,
    }
    validate_profile(result["profile"])
    return result


def archive_template(
    name: str, library_dir: Path = DEFAULT_LIBRARY_DIR
) -> dict[str, Any]:
    with library_lock(library_dir):
        registry = load_registry(library_dir)
        wanted = name_key(name)
        entry = next(
            (item for item in registry["active"] if item["name_key"] == wanted), None
        )
        if entry is None:
            raise StateError("未找到要归档的有效模板。")
        registry["active"].remove(entry)
        entry = {**entry, "archived_at": utc_now()}
        registry["archived"].append(entry)
        save_registry(registry, library_dir)
    return {"status": "archived", "template": entry, "remaining_slots": MAX_CUSTOM_TEMPLATES - len(registry["active"])}


def restore_template(
    name: str, library_dir: Path = DEFAULT_LIBRARY_DIR
) -> dict[str, Any]:
    with library_lock(library_dir):
        registry = load_registry(library_dir)
        if len(registry["active"]) >= MAX_CUSTOM_TEMPLATES:
            raise StateError("三槽已满，不能恢复模板；请先归档另一套。")
        wanted = name_key(name)
        entry = next(
            (item for item in registry["archived"] if item["name_key"] == wanted), None
        )
        if entry is None:
            raise StateError("未找到可恢复的归档模板。")
        if any(item["name_key"] == wanted for item in registry["active"]):
            raise StateError("已有同名有效模板。")
        registry["archived"].remove(entry)
        restored = {key: value for key, value in entry.items() if key != "archived_at"}
        registry["active"].append(restored)
        save_registry(registry, library_dir)
    return {"status": "restored", "template": restored, "remaining_slots": MAX_CUSTOM_TEMPLATES - len(registry["active"])}


def rename_template(
    current_name: str, new_name: str, library_dir: Path = DEFAULT_LIBRARY_DIR
) -> dict[str, Any]:
    cleaned = validate_template_name(new_name)
    with library_lock(library_dir):
        registry = load_registry(library_dir)
        wanted = name_key(current_name)
        entry = next(
            (item for item in registry["active"] if item["name_key"] == wanted), None
        )
        if entry is None:
            raise StateError("未找到要重命名的有效模板。")
        if any(
            item["name_key"] == name_key(cleaned)
            for item in registry["active"]
            if item["id"] != entry["id"]
        ):
            raise StateError("新名称已被其他模板占用。")
        entry["name"] = cleaned
        entry["name_key"] = name_key(cleaned)
        metadata_path = (
            library_dir / "templates" / entry["id"] / "metadata.json"
        )
        metadata = load_json(metadata_path)
        metadata["name"] = cleaned
        metadata["name_key"] = name_key(cleaned)
        atomic_write_json(metadata_path, metadata)
        profile_path = library_dir / "templates" / entry["id"] / "profile.json"
        profile = load_json(profile_path)
        profile["name"] = cleaned
        atomic_write_json(profile_path, profile)
        save_registry(registry, library_dir)
    return {"status": "renamed", "template": entry}


def show_template(
    name: str, library_dir: Path = DEFAULT_LIBRARY_DIR
) -> dict[str, Any]:
    template = get_template(name, library_dir)
    return {
        "id": template["id"],
        "name": template["name"],
        "mode": template["mode"],
        "builtin": template["builtin"],
        "profile": template["profile"],
        "shell_available": bool(template.get("shell_path")),
    }
