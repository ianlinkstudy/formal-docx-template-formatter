from __future__ import annotations

import os
import tempfile
from pathlib import Path
from typing import Any

from docx.document import Document as DocumentObject
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

from .common import (
    ANCHOR_RE,
    FormatterError,
    atomic_write_json,
    classify_roles,
    document_manifest,
    normalize_anchor,
    open_document,
    stable_hash,
)

ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "distribute": WD_ALIGN_PARAGRAPH.DISTRIBUTE,
}
SPECIAL_INLINE_MARKERS = (
    "<w:hyperlink",
    "<w:fldChar",
    "<w:drawing",
    "<w:object",
    "<m:oMath",
    "<w:commentRangeStart",
    "<w:bookmarkStart",
)


def _set_run_font(
    run: Any,
    font: str,
    latin_font: str,
    size_pt: float,
    bold: bool | None,
    italic: bool | None,
    color: str,
) -> None:
    run.font.name = latin_font
    run.font.size = Pt(float(size_pt))
    if bold is not None:
        run.font.bold = bool(bold)
    if italic is not None:
        run.font.italic = bool(italic)
    try:
        run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))
    except (TypeError, ValueError):
        pass
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        key = qn(f"w:{attr}")
        if key in r_fonts.attrib:
            del r_fonts.attrib[key]
    r_fonts.set(qn("w:ascii"), latin_font)
    r_fonts.set(qn("w:hAnsi"), latin_font)
    r_fonts.set(qn("w:cs"), latin_font)
    r_fonts.set(qn("w:eastAsia"), font)


def _format_runs(paragraph: Paragraph, style: dict[str, Any]) -> None:
    for run in paragraph.runs:
        _set_run_font(
            run,
            style["font"],
            style["latin_font"],
            style["size_pt"],
            style.get("bold"),
            style.get("italic"),
            style.get("color", "000000"),
        )


def _set_indent(
    paragraph: Paragraph,
    first_line_chars: float,
    right_chars: float,
    size_pt: float,
) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    ind = p_pr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        p_pr.append(ind)
    for attr in (
        "firstLine",
        "firstLineChars",
        "left",
        "leftChars",
        "hanging",
        "right",
        "rightChars",
        "end",
        "endChars",
    ):
        key = qn(f"w:{attr}")
        if key in ind.attrib:
            del ind.attrib[key]
    if first_line_chars > 0:
        ind.set(qn("w:firstLineChars"), str(round(first_line_chars * 100)))
        ind.set(
            qn("w:firstLine"),
            str(round(first_line_chars * size_pt * 20)),
        )
    if right_chars > 0:
        ind.set(qn("w:rightChars"), str(round(right_chars * 100)))
        ind.set(
            qn("w:right"),
            str(round(right_chars * size_pt * 20)),
        )


def _set_line_spacing(paragraph: Paragraph, style: dict[str, Any]) -> None:
    fmt = paragraph.paragraph_format
    rule = style.get("line_rule", "single")
    value = style.get("line_value")
    if rule in {"exact", "at_least"}:
        fmt.line_spacing = Pt(float(value))
        fmt.line_spacing_rule = (
            WD_LINE_SPACING.EXACTLY
            if rule == "exact"
            else WD_LINE_SPACING.AT_LEAST
        )
    elif rule == "multiple":
        fmt.line_spacing = float(value)
    else:
        fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.space_before = Pt(float(style.get("space_before_pt", 0)))
    fmt.space_after = Pt(float(style.get("space_after_pt", 0)))


def _apply_paragraph_style(
    paragraph: Paragraph,
    style: dict[str, Any],
    preserve_alignment: bool = False,
) -> None:
    fmt = paragraph.paragraph_format
    if not preserve_alignment:
        fmt.alignment = ALIGNMENTS.get(
            style.get("alignment", "justify"), WD_ALIGN_PARAGRAPH.JUSTIFY
        )
    _set_line_spacing(paragraph, style)
    _set_indent(
        paragraph,
        float(style.get("first_line_chars", 0)),
        float(style.get("right_chars", 0)),
        float(style.get("size_pt", 16)),
    )
    fmt.keep_with_next = bool(style.get("keep_with_next", False))


def _split_heading2(paragraph: Paragraph, profile: dict[str, Any]) -> tuple[bool, str | None]:
    text = paragraph.text
    index = text.find("。")
    if index < 0 or index == len(text) - 1:
        return False, None
    if any(marker in paragraph._p.xml for marker in SPECIAL_INLINE_MARKERS):
        return False, "二级标题与正文同段，但含超链接、域、图片或公式；整段保留为一个 run。"
    heading = text[: index + 1]
    body = text[index + 1 :]
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    heading_run = paragraph.add_run(heading)
    body_run = paragraph.add_run(body)
    _set_run_font(
        heading_run,
        **{
            key: profile["styles"]["heading2"][key]
            for key in ("font", "latin_font", "size_pt", "bold", "italic", "color")
        },
    )
    _set_run_font(
        body_run,
        **{
            key: profile["styles"]["body"][key]
            for key in ("font", "latin_font", "size_pt", "bold", "italic", "color")
        },
    )
    _apply_paragraph_style(paragraph, profile["styles"]["heading2"])
    return True, None


def _set_cell_shading(cell: Any, fill: str | None) -> None:
    if not fill:
        return
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill.replace("#", ""))
    shading.set(qn("w:val"), "clear")


def _set_table_borders(table: Table, style: dict[str, Any]) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(style["border_size_eighths_pt"]))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), style["border_color"])


def _set_table_margins(table: Table, style: dict[str, Any]) -> None:
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for edge in ("top", "left", "bottom", "right"):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(style[f"cell_margin_{edge}_twips"]))
        element.set(qn("w:type"), "dxa")


def _visual_units(text: str) -> int:
    return sum(2 if ord(character) > 127 else 1 for character in text.strip())


def _fit_table(table: Table, profile: dict[str, Any]) -> None:
    if not profile["table"].get("fit_to_usable_width", True):
        return
    columns = list(table._tbl.tblGrid.gridCol_lst)
    if not columns:
        return
    page = profile["page"]
    usable_mm = (
        float(page["width_mm"])
        - float(page["margin_left_mm"])
        - float(page["margin_right_mm"])
    )
    font_mm = float(profile["table"]["body_size_pt"]) * 25.4 / 72
    padding_mm = (
        float(profile["table"]["cell_margin_left_twips"])
        + float(profile["table"]["cell_margin_right_twips"])
    ) * 25.4 / 1440
    desired = [0.0] * len(columns)
    for row in table.rows:
        column_index = 0
        for tc in row._tr.tc_lst:
            tc_pr = tc.tcPr
            span_element = (
                tc_pr.find(qn("w:gridSpan")) if tc_pr is not None else None
            )
            span = (
                int(span_element.get(qn("w:val")))
                if span_element is not None
                else 1
            )
            if span == 1 and column_index < len(desired):
                text = "".join(node.text or "" for node in tc.findall(".//" + qn("w:t")))
                desired[column_index] = max(
                    desired[column_index],
                    _visual_units(text) * font_mm / 2 + padding_mm,
                )
            column_index += span
    minimum = min(
        float(profile["table"].get("minimum_column_mm", 12)),
        usable_mm / len(columns),
    )
    desired = [min(max(value, minimum), usable_mm * 0.45) for value in desired]
    total = sum(desired)
    widths_mm = (
        [usable_mm / len(columns)] * len(columns)
        if total <= 0
        else [value / total * usable_mm for value in desired]
    )
    widths = [max(1, round(value * 1440 / 25.4)) for value in widths_mm]
    target = round(usable_mm * 1440 / 25.4)
    widths[-1] += target - sum(widths)
    for column, width in zip(columns, widths):
        column.set(qn("w:w"), str(width))
    tbl_pr = table._tbl.tblPr
    tbl_width = tbl_pr.find(qn("w:tblW"))
    if tbl_width is None:
        tbl_width = OxmlElement("w:tblW")
        tbl_pr.append(tbl_width)
    tbl_width.set(qn("w:w"), str(target))
    tbl_width.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for row in table.rows:
        column_index = 0
        for tc in row._tr.tc_lst:
            tc_pr = tc.get_or_add_tcPr()
            span_element = tc_pr.find(qn("w:gridSpan"))
            span = (
                int(span_element.get(qn("w:val")))
                if span_element is not None
                else 1
            )
            width = sum(widths[column_index : column_index + span])
            tc_width = tc_pr.find(qn("w:tcW"))
            if tc_width is None:
                tc_width = OxmlElement("w:tcW")
                tc_pr.append(tc_width)
            tc_width.set(qn("w:w"), str(width))
            tc_width.set(qn("w:type"), "dxa")
            column_index += span


def _format_table(table: Table, profile: dict[str, Any]) -> None:
    table_style = profile["table"]
    _set_table_borders(table, table_style)
    _set_table_margins(table, table_style)
    _fit_table(table, profile)
    for row_index, row in enumerate(table.rows):
        is_header = row_index == 0
        if is_header and table_style.get("repeat_header_row"):
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:tblHeader")) is None:
                tr_pr.append(OxmlElement("w:tblHeader"))
        font = table_style["header_font"] if is_header else table_style["body_font"]
        size = (
            table_style["header_size_pt"]
            if is_header
            else table_style["body_size_pt"]
        )
        bold = (
            table_style["header_bold"] if is_header else table_style["body_bold"]
        )
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if is_header:
                _set_cell_shading(cell, table_style.get("header_fill"))
            for paragraph in cell.paragraphs:
                paragraph.alignment = ALIGNMENTS.get(
                    table_style.get("alignment", "center"),
                    WD_ALIGN_PARAGRAPH.CENTER,
                )
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                _set_indent(paragraph, 0, 0, float(size))
                run_style = {
                    "font": font,
                    "latin_font": table_style["latin_font"],
                    "size_pt": size,
                    "bold": bold,
                    "italic": None,
                    "color": "000000",
                }
                _format_runs(paragraph, run_style)


def _resize_inline_images(doc: DocumentObject, profile: dict[str, Any]) -> int:
    if not doc.sections:
        return 0
    usable_widths = []
    for section in doc.sections:
        usable_widths.append(
            section.page_width
            - section.left_margin
            - section.right_margin
        )
    max_width = int(
        int(min(usable_widths))
        * float(profile["figure"].get("max_width_ratio", 1.0))
    )
    resized = 0
    for shape in doc.inline_shapes:
        if int(shape.width) <= max_width:
            continue
        ratio = max_width / int(shape.width)
        old_height = int(shape.height)
        shape.width = max_width
        shape.height = max(1, round(old_height * ratio))
        resized += 1
    for paragraph in doc.paragraphs:
        if "<w:drawing" not in paragraph._p.xml or paragraph.text.strip():
            continue
        paragraph.alignment = ALIGNMENTS.get(
            profile["figure"].get("inline_alignment", "center"),
            WD_ALIGN_PARAGRAPH.CENTER,
        )
        paragraph.paragraph_format.first_line_indent = Pt(0)
    return resized


def format_style_document(
    doc: DocumentObject, profile: dict[str, Any]
) -> dict[str, Any]:
    before = document_manifest(doc)
    roles = classify_roles(doc)
    warnings: list[str] = []
    split_count = 0
    formatted_paragraphs = 0

    page = profile["page"]
    for section in doc.sections:
        section.page_width = Mm(float(page["width_mm"]))
        section.page_height = Mm(float(page["height_mm"]))
        section.orientation = (
            WD_ORIENT.LANDSCAPE
            if float(page["width_mm"]) > float(page["height_mm"])
            else WD_ORIENT.PORTRAIT
        )
        section.top_margin = Mm(float(page["margin_top_mm"]))
        section.bottom_margin = Mm(float(page["margin_bottom_mm"]))
        section.left_margin = Mm(float(page["margin_left_mm"]))
        section.right_margin = Mm(float(page["margin_right_mm"]))

    for index, paragraph in enumerate(doc.paragraphs):
        if not paragraph.text.strip():
            continue
        role = roles[index]
        style = profile["styles"].get(role, profile["styles"]["body"])
        if (
            "<w:drawing" in paragraph._p.xml
            and not paragraph.text.strip()
        ):
            style = {**style, "alignment": profile["figure"].get("inline_alignment", "center")}
            style["first_line_chars"] = 0
        preserve_alignment = bool(style.get("preserve_existing_alignment"))
        if role == "heading2":
            split, warning = _split_heading2(paragraph, profile)
            if split:
                split_count += 1
            elif warning:
                warnings.append(f"第{index + 1}段：{warning}")
                _format_runs(paragraph, style)
                _apply_paragraph_style(paragraph, style)
            else:
                _format_runs(paragraph, style)
                _apply_paragraph_style(paragraph, style)
        else:
            _format_runs(paragraph, style)
            _apply_paragraph_style(
                paragraph, style, preserve_alignment=preserve_alignment
            )
        formatted_paragraphs += 1

    for table in doc.tables:
        _format_table(table, profile)

    nested_tables = doc.element.xml.count("<w:tbl>") - len(doc.tables)
    if nested_tables > 0:
        warnings.append(f"检测到 {nested_tables} 个嵌套表格，已保留原样。")
    floating = doc.element.xml.count("<wp:anchor")
    if floating:
        warnings.append(f"检测到 {floating} 个浮动对象，已保持原位。")
    ole_count = doc.element.xml.count("<w:object")
    if ole_count:
        warnings.append(f"检测到 {ole_count} 个 OLE/MathType 对象，已原样保留并核验。")
    resized_images = _resize_inline_images(doc, profile)

    after = document_manifest(doc)
    if before != after:
        raise FormatterError("排版后正文、表格结构、公式、图片或页眉页脚发生变化，已停止保存。")
    return {
        "before": before,
        "after": after,
        "semantic_hash": stable_hash(after),
        "formatted_paragraphs": formatted_paragraphs,
        "formatted_tables": len(doc.tables),
        "split_heading2": split_count,
        "resized_inline_images": resized_images,
        "formula_count": len(after["formula_hashes"]),
        "ole_object_count": len(after["object_hashes"]),
        "nested_table_count": max(0, nested_tables),
        "floating_object_count": floating,
        "warnings": warnings,
    }


def _save_doc_atomic(doc: DocumentObject, target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(
        prefix=f".{target.stem}.", suffix=".docx", dir=target.parent
    )
    os.close(fd)
    try:
        doc.save(temp_name)
        os.replace(temp_name, target)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)


def format_style_file(
    input_path: Path,
    output_path: Path,
    template: dict[str, Any],
    overwrite: bool = False,
) -> dict[str, Any]:
    if input_path.resolve() == output_path.resolve():
        raise FormatterError("输出路径必须与输入路径不同。")
    if output_path.exists() and not overwrite:
        raise FormatterError("输出文件已存在；如需覆盖请显式启用 overwrite。")
    if template["mode"] != "style":
        raise FormatterError("该模板是母版模式，请使用母版排版流程。")
    doc = open_document(input_path)
    result = format_style_document(doc, template["profile"])
    _save_doc_atomic(doc, output_path)
    return {
        **result,
        "template": template["name"],
        "template_id": template["id"],
        "template_mode": "style",
        "output": str(output_path),
    }


def _source_role_blocks(doc: DocumentObject) -> tuple[list[Any], dict[str, list[int]]]:
    blocks = [
        child
        for child in doc.element.body.iterchildren()
        if child.tag in {qn("w:p"), qn("w:tbl")}
    ]
    roles = classify_roles(doc)
    paragraph_to_role = {
        paragraph._p: roles[index] for index, paragraph in enumerate(doc.paragraphs)
    }
    grouped = {
        "TITLE": [],
        "RECIPIENT": [],
        "SIGNATURE": [],
        "DATE": [],
    }
    for index, block in enumerate(blocks):
        if block.tag != qn("w:p"):
            continue
        role = paragraph_to_role.get(block)
        if role == "title":
            grouped["TITLE"].append(index)
        elif role == "recipient":
            grouped["RECIPIENT"].append(index)
        elif role == "signature":
            grouped["SIGNATURE"].append(index)
        elif role == "date":
            grouped["DATE"].append(index)
    return blocks, grouped


def _filtered_source(source: Path, selected: set[int]) -> DocumentObject:
    doc = open_document(source)
    blocks = [
        child
        for child in doc.element.body.iterchildren()
        if child.tag in {qn("w:p"), qn("w:tbl")}
    ]
    for index, block in enumerate(blocks):
        if index not in selected and block.getparent() is not None:
            block.getparent().remove(block)
    return doc


def _marker_paragraphs(doc: DocumentObject) -> dict[str, Paragraph]:
    result: dict[str, Paragraph] = {}
    for paragraph in doc.paragraphs:
        for match in ANCHOR_RE.finditer(paragraph.text):
            canonical = normalize_anchor(match.group(1))
            if canonical:
                result[canonical] = paragraph
    return result


def _is_subset_manifest(source: dict[str, Any], output: dict[str, Any]) -> bool:
    source_paragraphs = [
        block["text"]
        for block in source["blocks"]
        if block["kind"] == "paragraph" and block["text"].strip()
    ]
    output_paragraphs = [
        block["text"]
        for block in output["blocks"]
        if block["kind"] == "paragraph" and block["text"].strip()
    ]
    source_tables = [
        block["signature"] for block in source["blocks"] if block["kind"] == "table"
    ]
    output_tables = [
        block["signature"] for block in output["blocks"] if block["kind"] == "table"
    ]
    from collections import Counter

    paragraph_ok = not (Counter(source_paragraphs) - Counter(output_paragraphs))
    table_hashes = Counter(stable_hash(item) for item in source_tables)
    output_table_hashes = Counter(stable_hash(item) for item in output_tables)
    table_ok = not (table_hashes - output_table_hashes)
    formula_ok = set(source["formula_hashes"]).issubset(output["formula_hashes"])
    object_ok = set(source["object_hashes"]).issubset(output["object_hashes"])
    image_ok = set(source["part_hashes"]["images"]).issubset(
        output["part_hashes"]["images"]
    )
    embedding_ok = set(source["part_hashes"]["embeddings"]).issubset(
        output["part_hashes"]["embeddings"]
    )
    return all(
        (paragraph_ok, table_ok, formula_ok, object_ok, image_ok, embedding_ok)
    )


def _template_fixed_retained(shell: dict[str, Any], output: dict[str, Any]) -> bool:
    from collections import Counter

    fixed_paragraphs = [
        block["text"]
        for block in shell["blocks"]
        if block["kind"] == "paragraph"
        and block["text"].strip()
        and not ANCHOR_RE.search(block["text"])
    ]
    output_paragraphs = [
        block["text"]
        for block in output["blocks"]
        if block["kind"] == "paragraph" and block["text"].strip()
    ]
    fixed_tables = Counter(
        stable_hash(block["signature"])
        for block in shell["blocks"]
        if block["kind"] == "table"
    )
    output_tables = Counter(
        stable_hash(block["signature"])
        for block in output["blocks"]
        if block["kind"] == "table"
    )
    return all(
        (
            not (Counter(fixed_paragraphs) - Counter(output_paragraphs)),
            not (fixed_tables - output_tables),
            shell["headers_footers"] == output["headers_footers"],
            set(shell["formula_hashes"]).issubset(set(output["formula_hashes"])),
            set(shell["object_hashes"]).issubset(set(output["object_hashes"])),
            set(shell["part_hashes"]["images"]).issubset(
                set(output["part_hashes"]["images"])
            ),
        )
    )


def format_shell_file(
    input_path: Path,
    output_path: Path,
    template: dict[str, Any],
    overwrite: bool = False,
) -> dict[str, Any]:
    if input_path.resolve() == output_path.resolve():
        raise FormatterError("输出路径必须与输入路径不同。")
    if output_path.exists() and not overwrite:
        raise FormatterError("输出文件已存在；如需覆盖请显式启用 overwrite。")
    if template["mode"] != "shell" or not template.get("shell_path"):
        raise FormatterError("母版文件不存在或模板不是母版模式。")
    source_doc = open_document(input_path)
    if len(source_doc.sections) != 1:
        raise FormatterError("母版模式首版只接受单节输入正文；请改用样式模式。")
    try:
        from docxcompose.composer import Composer
    except ImportError as exc:
        raise FormatterError("缺少 docxcompose。请安装 requirements.txt。") from exc

    with tempfile.TemporaryDirectory() as temp_dir:
        temp = Path(temp_dir)
        formatted_source = temp / "source-formatted.docx"
        formatting_doc = open_document(input_path)
        source_style = format_style_document(formatting_doc, template["profile"])
        _save_doc_atomic(formatting_doc, formatted_source)
        source_manifest = document_manifest(open_document(formatted_source))
        working_doc = open_document(formatted_source)
        blocks, grouped = _source_role_blocks(working_doc)
        markers = _marker_paragraphs(open_document(template["shell_path"]))
        selected: dict[str, set[int]] = {}
        used: set[int] = set()
        for role in ("TITLE", "RECIPIENT", "SIGNATURE", "DATE"):
            if role in markers:
                selected[role] = set(grouped[role])
                used.update(grouped[role])
        selected["BODY"] = {
            index for index in range(len(blocks)) if index not in used
        }

        shell_doc = open_document(template["shell_path"])
        markers = _marker_paragraphs(shell_doc)
        composer = Composer(shell_doc)
        for role in sorted(
            markers,
            key=lambda item: shell_doc.element.body.index(markers[item]._p),
            reverse=True,
        ):
            marker_paragraph = markers[role]
            insertion_index = shell_doc.element.body.index(marker_paragraph._p)
            indices = selected.get(role, set())
            if indices:
                filtered = _filtered_source(formatted_source, indices)
                composer.insert(insertion_index, filtered)
            if marker_paragraph._p.getparent() is not None:
                marker_paragraph._p.getparent().remove(marker_paragraph._p)
        _save_doc_atomic(shell_doc, output_path)

    shell_manifest = document_manifest(open_document(template["shell_path"]))
    output_manifest = document_manifest(open_document(output_path))
    if not _is_subset_manifest(source_manifest, output_manifest):
        raise FormatterError("母版合并后未通过源正文完整性核验。")
    if not _template_fixed_retained(shell_manifest, output_manifest):
        raise FormatterError("母版固定内容、页眉页脚或对象在合并后发生变化。")
    warnings = list(source_style["warnings"])
    return {
        **source_style,
        "semantic_hash": source_style["semantic_hash"],
        "template": template["name"],
        "template_id": template["id"],
        "template_mode": "shell",
        "output": str(output_path),
        "source_manifest": source_manifest,
        "output_manifest": output_manifest,
        "template_embedded_formula_count": len(shell_manifest["formula_hashes"]),
        "warnings": warnings,
    }


def format_with_template(
    input_path: Path,
    output_path: Path,
    template: dict[str, Any],
    overwrite: bool = False,
) -> dict[str, Any]:
    if template["mode"] == "shell":
        return format_shell_file(input_path, output_path, template, overwrite)
    return format_style_file(input_path, output_path, template, overwrite)


def _actual_first_run_style(paragraph: Paragraph) -> dict[str, Any] | None:
    run = next((item for item in paragraph.runs if item.text.strip()), None)
    if run is None:
        return None
    r_pr = run._element.rPr
    fonts = r_pr.rFonts if r_pr is not None else None
    return {
        "font": fonts.get(qn("w:eastAsia")) if fonts is not None else None,
        "size_pt": round(run.font.size.pt, 2) if run.font.size is not None else None,
        "bold": run.bold,
        "alignment": paragraph.paragraph_format.alignment,
    }


def _indent_value(paragraph: Paragraph, attribute: str) -> int | None:
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return None
    ind = p_pr.find(qn("w:ind"))
    if ind is None:
        return None
    value = ind.get(qn(f"w:{attribute}"))
    return int(value) if value is not None else 0


def _line_spacing_matches(paragraph: Paragraph, expected: dict[str, Any]) -> bool:
    spacing = paragraph.paragraph_format.line_spacing
    rule = paragraph.paragraph_format.line_spacing_rule
    wanted_rule = expected.get("line_rule", "single")
    wanted_value = expected.get("line_value")
    if wanted_rule == "multiple":
        return isinstance(spacing, (int, float)) and abs(
            float(spacing) - float(wanted_value)
        ) <= 0.02
    if wanted_rule in {"exact", "at_least"}:
        if not hasattr(spacing, "pt"):
            return False
        wanted_enum = (
            WD_LINE_SPACING.EXACTLY
            if wanted_rule == "exact"
            else WD_LINE_SPACING.AT_LEAST
        )
        return (
            rule == wanted_enum
            and abs(float(spacing.pt) - float(wanted_value)) <= 0.2
        )
    return rule == WD_LINE_SPACING.SINGLE


def verify_file(
    formatted_path: Path,
    reference_path: Path,
    template: dict[str, Any],
) -> dict[str, Any]:
    formatted = open_document(formatted_path)
    reference = open_document(reference_path)
    checks: list[dict[str, Any]] = []
    output_manifest = document_manifest(formatted)
    reference_manifest = document_manifest(reference)
    if template["mode"] == "style":
        semantic_ok = output_manifest == reference_manifest
        checks.append(
            {
                "name": "正文与富对象完整性",
                "passed": semantic_ok,
                "evidence": (
                    "文本、表格、页眉页脚、图片与公式哈希一致"
                    if semantic_ok
                    else "检测到正文或对象变化"
                ),
            }
        )
    else:
        semantic_ok = _is_subset_manifest(reference_manifest, output_manifest)
        shell_manifest = document_manifest(open_document(template["shell_path"]))
        fixed_ok = _template_fixed_retained(shell_manifest, output_manifest)
        checks.append(
            {
                "name": "源正文已全部迁入母版",
                "passed": semantic_ok,
                "evidence": "源文本、表格、图片、公式哈希均可在输出中找到",
            }
        )
        checks.append(
            {
                "name": "母版页眉页脚保留",
                "passed": fixed_ok,
                "evidence": "母版固定文本、对象与页眉页脚结构一致",
            }
        )

    page = template["profile"]["page"]
    page_ok = True
    for section in formatted.sections:
        values = (
            round(section.page_width.mm, 1),
            round(section.page_height.mm, 1),
            round(section.top_margin.mm, 1),
            round(section.bottom_margin.mm, 1),
            round(section.left_margin.mm, 1),
            round(section.right_margin.mm, 1),
        )
        expected = (
            round(float(page["width_mm"]), 1),
            round(float(page["height_mm"]), 1),
            round(float(page["margin_top_mm"]), 1),
            round(float(page["margin_bottom_mm"]), 1),
            round(float(page["margin_left_mm"]), 1),
            round(float(page["margin_right_mm"]), 1),
        )
        if template["mode"] == "style" and any(
            abs(actual - wanted) > 0.3 for actual, wanted in zip(values, expected)
        ):
            page_ok = False
    checks.append(
        {
            "name": "页面与页边距",
            "passed": page_ok,
            "evidence": "样式模式按所选模板页面参数核验；母版模式保留母版分节",
        }
    )

    typography_failures: list[str] = []
    layout_failures: list[str] = []
    if template["mode"] == "style":
        roles = classify_roles(formatted)
        for index, paragraph in enumerate(formatted.paragraphs):
            if not paragraph.text.strip():
                continue
            role = roles[index]
            expected = template["profile"]["styles"].get(
                role, template["profile"]["styles"]["body"]
            )
            if "<m:oMath" not in paragraph._p.xml:
                actual = _actual_first_run_style(paragraph)
                if actual is not None:
                    if actual["font"] != expected["font"]:
                        typography_failures.append(
                            f"第{index + 1}段字体={actual['font']}，应为 {expected['font']}"
                        )
                    if actual["size_pt"] != round(float(expected["size_pt"]), 2):
                        typography_failures.append(
                            f"第{index + 1}段字号={actual['size_pt']}，应为 {expected['size_pt']}"
                        )
                    if expected.get("bold") is not None and bool(actual["bold"]) != bool(
                        expected["bold"]
                    ):
                        typography_failures.append(
                            f"第{index + 1}段字重={actual['bold']}，应为 {expected['bold']}"
                        )
            if not expected.get("preserve_existing_alignment"):
                wanted_alignment = ALIGNMENTS.get(
                    expected.get("alignment", "justify"),
                    WD_ALIGN_PARAGRAPH.JUSTIFY,
                )
                if paragraph.paragraph_format.alignment != wanted_alignment:
                    layout_failures.append(
                        f"第{index + 1}段对齐={paragraph.paragraph_format.alignment}"
                    )
            if not _line_spacing_matches(paragraph, expected):
                layout_failures.append(f"第{index + 1}段行距不符合模板")
            expected_first = round(
                float(expected.get("first_line_chars", 0)) * 100
            )
            expected_right = round(
                float(expected.get("right_chars", 0)) * 100
            )
            if (
                (_indent_value(paragraph, "firstLineChars") or 0) != expected_first
                or (_indent_value(paragraph, "rightChars") or 0) != expected_right
            ):
                layout_failures.append(f"第{index + 1}段缩进不符合模板")
    checks.append(
        {
            "name": "段落字体与字号",
            "passed": not typography_failures,
            "evidence": (
                "全部通过"
                if not typography_failures
                else "；".join(typography_failures[:8])
            ),
        }
    )

    checks.append(
        {
            "name": "段落对齐、行距与缩进",
            "passed": not layout_failures,
            "evidence": (
                "全部通过"
                if not layout_failures
                else "；".join(layout_failures[:8])
            ),
        }
    )

    table_failures: list[str] = []
    table_style = template["profile"]["table"]
    for table_index, table in enumerate(formatted.tables):
        for row_index, row in enumerate(table.rows):
            expected_font = (
                table_style["header_font"]
                if row_index == 0
                else table_style["body_font"]
            )
            expected_size = (
                table_style["header_size_pt"]
                if row_index == 0
                else table_style["body_size_pt"]
            )
            for cell_index, cell in enumerate(row.cells):
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if not run.text.strip():
                            continue
                        r_pr = run._element.rPr
                        fonts = r_pr.rFonts if r_pr is not None else None
                        actual_font = (
                            fonts.get(qn("w:eastAsia")) if fonts is not None else None
                        )
                        actual_size = (
                            round(run.font.size.pt, 2)
                            if run.font.size is not None
                            else None
                        )
                        if actual_font != expected_font or actual_size != round(
                            float(expected_size), 2
                        ):
                            table_failures.append(
                                f"表{table_index + 1} 行{row_index + 1} 单元{cell_index + 1}"
                            )
                            break
    checks.append(
        {
            "name": "表格字体与字号",
            "passed": not table_failures,
            "evidence": (
                "全部通过"
                if not table_failures
                else "；".join(table_failures[:8])
            ),
        }
    )
    passed = all(check["passed"] for check in checks)
    return {
        "schema_version": 1,
        "status": "passed" if passed else "failed",
        "formatted": str(formatted_path),
        "reference": str(reference_path),
        "template": template["name"],
        "template_mode": template["mode"],
        "checks": checks,
        "rich_objects": {
            "formula_count": len(output_manifest["formula_hashes"]),
            "ole_object_count": len(output_manifest["object_hashes"]),
            "inline_image_count": output_manifest["inline_image_count"],
            "floating_object_count": output_manifest["floating_drawing_count"],
            "image_part_hashes": output_manifest["part_hashes"]["images"],
        },
        "warnings": [],
    }


def write_report(path: Path, payload: dict[str, Any]) -> None:
    atomic_write_json(path, payload)
