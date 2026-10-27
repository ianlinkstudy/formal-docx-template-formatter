from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from base64 import b64decode
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

from core.common import FormatterError, StateError, document_manifest
from core.formatter_core import (
    format_style_document,
    format_with_template,
    verify_file,
)
from core.template_store import (
    amend_candidate,
    archive_template,
    confirm_candidate,
    get_template,
    inspect_template,
    restore_template,
    save_candidate,
    template_choices,
)

PNG_BYTES = b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "/A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def add_formula(paragraph, expression: str = "x+y") -> None:
    paragraph._p.append(
        parse_xml(
            "<m:oMath "
            + nsdecls("m", "w")
            + "><m:r><m:t>"
            + expression
            + "</m:t></m:r></m:oMath>"
        )
    )


def make_template(path: Path) -> None:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "模板页眉"
    doc.sections[0].footer.paragraphs[0].text = "模板页脚"
    doc.add_paragraph("关于测试工作的报告")
    doc.add_paragraph("测试集团：")
    doc.add_paragraph("一、总体情况")
    paragraph = doc.add_paragraph("（一）工作进展。本季度完成率达到95%。")
    paragraph.alignment = 3
    table = doc.add_table(rows=2, cols=3)
    for row_index, values in enumerate(
        (("序号", "事项", "状态"), ("1", "测试任务", "完成"))
    ):
        for column_index, value in enumerate(values):
            table.cell(row_index, column_index).text = value
    table.cell(1, 1).merge(table.cell(1, 2))
    doc.add_paragraph("特此报告。")
    doc.add_paragraph("测试集团有限公司").alignment = 2
    doc.add_paragraph("2026年10月27日").alignment = 2
    for item in doc.paragraphs:
        for run in item.runs:
            run.font.size = Pt(16)
    doc.save(path)


def make_shell_template(path: Path) -> None:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "公司固定页眉"
    doc.sections[0].footer.paragraphs[0].text = "公司固定页脚"
    doc.add_paragraph("公司固定封面")
    doc.add_paragraph("{{TITLE}}")
    doc.add_paragraph("{{RECIPIENT}}")
    doc.add_paragraph("{{BODY}}")
    doc.add_paragraph("{{SIGNATURE}}")
    doc.add_paragraph("{{DATE}}")
    doc.save(path)


def make_source(path: Path, include_rich: bool = False) -> None:
    doc = Document()
    doc.add_paragraph("关于测试事项的报告")
    doc.add_paragraph("测试集团：")
    doc.add_paragraph("一、总体情况")
    doc.add_paragraph("这是原始正文，不得改动。")
    if include_rich:
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "项目"
        table.cell(0, 1).text = "结果"
        table.cell(1, 0).text = "A"
        table.cell(1, 1).text = "完成"
        doc.add_picture(BytesIO(PNG_BYTES))
        image = doc.inline_shapes[-1]
        image.width = Inches(10)
        image.height = Inches(1)
        formula = doc.add_paragraph()
        add_formula(formula)
        ole = doc.add_paragraph()
        ole._p.append(
            parse_xml(
                "<w:object "
                + nsdecls("w")
                + '><w:shape id="Object 1"/></w:object>'
            )
        )
    doc.add_paragraph("特此报告。")
    doc.add_paragraph("测试集团有限公司")
    doc.add_paragraph("2026年10月27日")
    doc.save(path)


def inspect_confirm_save(
    source: Path, name: str, library: Path, mode: str = "auto"
) -> str:
    result = inspect_template(source, mode=mode, library_dir=library)
    candidate = result["candidate"]
    confirm_candidate(
        candidate["candidate_id"],
        candidate["candidate_sha256"],
        "确认解析结果",
        library_dir=library,
    )
    saved = save_candidate(
        candidate["candidate_id"], name, library_dir=library
    )
    return saved["template"]["id"]


class TemplateStateTests(unittest.TestCase):
    def test_first_selection_is_gb_only_and_requires_choice(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            choices = template_choices(Path(temp_dir) / "library")
            self.assertEqual(len(choices["builtin"]), 1)
            self.assertEqual(choices["builtin"][0]["id"], "gb9704")
            self.assertEqual(choices["custom"], [])
            self.assertEqual(choices["remaining_slots"], 3)
            self.assertTrue(choices["requires_selection"])

    def test_unconfirmed_candidate_cannot_be_saved_or_formatted(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            library = temp / "library"
            source = temp / "template.docx"
            make_template(source)
            result = inspect_template(source, "style", library_dir=library)
            candidate = result["candidate"]
            self.assertEqual(result["status"], "needs_confirmation")
            self.assertTrue(candidate["candidate_sha256"])
            with self.assertRaises(StateError):
                confirm_candidate(
                    candidate["candidate_id"],
                    candidate["candidate_sha256"],
                    "先这样",
                    library_dir=library,
                )
            with self.assertRaises(StateError):
                save_candidate(
                    candidate["candidate_id"], "未确认模板", library_dir=library
                )
            with self.assertRaises(StateError):
                get_template("未确认模板", library)

    def test_confirm_then_name_persists_but_amendment_invalidates_confirmation(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            library = temp / "library"
            source = temp / "template.docx"
            make_template(source)
            result = inspect_template(source, "style", library_dir=library)
            candidate = result["candidate"]
            confirm_candidate(
                candidate["candidate_id"],
                candidate["candidate_sha256"],
                "确认解析结果",
                library_dir=library,
            )
            with self.assertRaises(StateError):
                amend_candidate(
                    candidate["candidate_id"],
                    changes={"profile": {"styles": {"body": {"size_pt": 100}}}},
                    library_dir=library,
                )
            amended = amend_candidate(
                candidate["candidate_id"],
                changes={"profile": {"styles": {"body": {"line_value": 30}}}},
                library_dir=library,
            )
            self.assertNotEqual(
                amended["candidate"]["candidate_sha256"],
                candidate["candidate_sha256"],
            )
            with self.assertRaises(StateError):
                save_candidate(
                    candidate["candidate_id"],
                    "已失效模板",
                    library_dir=library,
                )
            with self.assertRaises(StateError):
                confirm_candidate(
                    candidate["candidate_id"],
                    candidate["candidate_sha256"],
                    "确认解析结果",
                    library_dir=library,
                )
            confirm_candidate(
                candidate["candidate_id"],
                amended["candidate"]["candidate_sha256"],
                "确认解析结果",
                library_dir=library,
            )
            saved = save_candidate(
                candidate["candidate_id"], "已确认模板", library_dir=library
            )
            self.assertEqual(saved["remaining_slots"], 2)
            self.assertEqual(
                len(template_choices(library)["custom"]),
                1,
            )

    def test_three_slot_limit_archive_and_restore(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            library = temp / "library"
            source = temp / "template.docx"
            make_template(source)
            for index in range(3):
                inspect_confirm_save(source, f"模板{index + 1}", library)
            self.assertEqual(template_choices(library)["remaining_slots"], 0)
            with self.assertRaises(FormatterError):
                inspect_template(source, "style", library_dir=library)
            archive_template("模板1", library)
            self.assertEqual(template_choices(library)["remaining_slots"], 1)
            inspect_confirm_save(source, "模板4", library)
            self.assertEqual(template_choices(library)["remaining_slots"], 0)
            with self.assertRaises(StateError):
                restore_template("模板1", library)

    def test_duplicate_name_rejected_and_style_source_removed_after_save(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            library = temp / "library"
            source = temp / "template.docx"
            make_template(source)
            result = inspect_template(source, "style", library_dir=library)
            candidate = result["candidate"]
            candidate_dir = library / "pending" / candidate["candidate_id"]
            self.assertTrue((candidate_dir / "source.docx").exists())
            confirm_candidate(
                candidate["candidate_id"],
                candidate["candidate_sha256"],
                "确认解析结果",
                library_dir=library,
            )
            save_candidate(candidate["candidate_id"], "集团模板", library_dir=library)
            self.assertFalse(candidate_dir.exists())
            second = inspect_template(source, "style", library_dir=library)["candidate"]
            confirm_candidate(
                second["candidate_id"],
                second["candidate_sha256"],
                "确认解析结果",
                library_dir=library,
            )
            with self.assertRaises(StateError):
                save_candidate(second["candidate_id"], "集团模板", library_dir=library)

    def test_dotx_is_accepted_without_modifying_source(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            library = temp / "library"
            docx_path = temp / "template.docx"
            dotx_path = temp / "template.dotx"
            make_template(docx_path)
            with ZipFile(docx_path) as zin, ZipFile(
                dotx_path, "w", ZIP_DEFLATED
            ) as zout:
                for item in zin.infolist():
                    payload = zin.read(item.filename)
                    if item.filename == "[Content_Types].xml":
                        payload = payload.replace(
                            b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                            b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
                        )
                    zout.writestr(item, payload)
            before = dotx_path.read_bytes()
            inspect_confirm_save(dotx_path, "DOTX模板", library)
            self.assertEqual(before, dotx_path.read_bytes())

    def test_style_inheritance_is_extracted(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            library = temp / "library"
            source = temp / "template.docx"
            doc = Document()
            doc.styles["Normal"].font.name = "Arial"
            doc.styles["Normal"].font.size = Pt(14)
            doc.styles["Heading 1"].font.name = "SimHei"
            doc.styles["Heading 1"].font.size = Pt(18)
            doc.add_paragraph("关于样式继承的报告")
            doc.add_paragraph("一、总体情况", style="Heading 1")
            doc.add_paragraph("这是使用 Normal 样式继承的正文。")
            doc.save(source)
            candidate = inspect_template(
                source, "style", library_dir=library
            )["candidate"]
            self.assertEqual(
                candidate["profile"]["styles"]["heading1"]["font"], "SimHei"
            )
            self.assertEqual(
                candidate["profile"]["styles"]["heading1"]["size_pt"], 18.0
            )
            self.assertEqual(
                candidate["profile"]["styles"]["body"]["size_pt"], 14.0
            )


class FormattingTests(unittest.TestCase):
    def test_style_format_preserves_formula_and_detects_tampering(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            source = temp / "source.docx"
            output = temp / "output.docx"
            tampered = temp / "tampered.docx"
            make_source(source, include_rich=True)
            template = get_template("gb9704", temp / "unused")
            result = format_with_template(
                source, output, template, overwrite=True
            )
            self.assertGreaterEqual(result["formula_count"], 1)
            self.assertGreaterEqual(result["ole_object_count"], 1)
            self.assertGreaterEqual(result["resized_inline_images"], 1)
            verification = verify_file(output, source, template)
            self.assertEqual(verification["status"], "passed")
            with ZipFile(output) as zin, ZipFile(
                tampered, "w", ZIP_DEFLATED
            ) as zout:
                for item in zin.infolist():
                    payload = zin.read(item.filename)
                    if item.filename == "word/document.xml":
                        payload = payload.replace(b"x+y", b"z-q")
                    zout.writestr(item, payload)
            tampered_result = verify_file(tampered, source, template)
            self.assertEqual(tampered_result["status"], "failed")
            self.assertFalse(tampered_result["checks"][0]["passed"])

    def test_oversized_inline_image_is_scaled_and_floating_untouched(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            source = temp / "source.docx"
            doc = Document()
            doc.add_paragraph("关于图片的报告")
            doc.add_picture(BytesIO(PNG_BYTES))
            doc.inline_shapes[0].width = Inches(20)
            doc.inline_shapes[0].height = Inches(2)
            doc.save(source)
            template = get_template("gb9704", temp / "unused")
            formatted = Document(source)
            result = format_style_document(formatted, template["profile"])
            self.assertEqual(result["resized_inline_images"], 1)
            usable_width = (
                formatted.sections[0].page_width
                - formatted.sections[0].left_margin
                - formatted.sections[0].right_margin
            )
            self.assertLessEqual(int(formatted.inline_shapes[0].width), int(usable_width))
            ratio_before = 2 / 20
            ratio_after = int(formatted.inline_shapes[0].height) / int(
                formatted.inline_shapes[0].width
            )
            self.assertAlmostEqual(ratio_before, ratio_after, places=2)

    def test_shell_preserves_fixed_areas_and_merges_rich_source(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            library = temp / "library"
            shell = temp / "shell.docx"
            source = temp / "source.docx"
            output = temp / "output.docx"
            make_shell_template(shell)
            make_source(source, include_rich=True)
            inspect_confirm_save(shell, "公司母版", library)
            template = get_template("公司母版", library)
            result = format_with_template(
                source, output, template, overwrite=True
            )
            output_doc = Document(output)
            texts = [paragraph.text for paragraph in output_doc.paragraphs]
            self.assertIn("公司固定封面", texts)
            self.assertIn("关于测试事项的报告", texts)
            self.assertIn("2026年10月27日", texts)
            self.assertEqual(output_doc.sections[0].header.paragraphs[0].text, "公司固定页眉")
            self.assertEqual(output_doc.sections[0].footer.paragraphs[0].text, "公司固定页脚")
            self.assertEqual(len(output_doc.tables), 1)
            self.assertEqual(len(output_doc.inline_shapes), 1)
            self.assertEqual(output_doc.element.xml.count("<m:oMath"), 1)
            self.assertEqual(output_doc.element.xml.count("<w:object"), 1)
            self.assertEqual(result["template_mode"], "shell")
            self.assertEqual(verify_file(output, source, template)["status"], "passed")

    def test_unknown_or_duplicate_shell_anchor_blocks_confirmation(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            library = temp / "library"
            source = temp / "shell.docx"
            doc = Document()
            doc.add_paragraph("{{BODY}}")
            doc.add_paragraph("{{BODY}}")
            doc.add_paragraph("{{UNKNOWN}}")
            doc.save(source)
            result = inspect_template(source, "shell", library_dir=library)
            candidate = result["candidate"]
            self.assertTrue(candidate["blockers"])
            with self.assertRaises(StateError):
                confirm_candidate(
                    candidate["candidate_id"],
                    candidate["candidate_sha256"],
                    "确认解析结果",
                    library_dir=library,
                )

    def test_header_anchor_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            library = temp / "library"
            source = temp / "shell.docx"
            doc = Document()
            doc.sections[0].header.paragraphs[0].text = "{{TITLE}}"
            doc.add_paragraph("{{BODY}}")
            doc.save(source)
            candidate = inspect_template(source, "shell", library_dir=library)["candidate"]
            self.assertTrue(
                any("页眉" in blocker or "表格" in blocker for blocker in candidate["blockers"])
            )
            with self.assertRaises(StateError):
                confirm_candidate(
                    candidate["candidate_id"],
                    candidate["candidate_sha256"],
                    "确认解析结果",
                    library_dir=library,
                )

    def test_explicit_body_range_materializes_shell_only_after_confirmation(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            library = temp / "library"
            shell = temp / "shell.docx"
            source = temp / "source.docx"
            output = temp / "output.docx"
            template_doc = Document()
            template_doc.add_paragraph("固定封面")
            template_doc.add_paragraph("示例标题")
            template_doc.add_paragraph("示例正文")
            template_doc.save(shell)
            result = inspect_template(
                shell,
                "shell",
                body_start=1,
                body_end=2,
                library_dir=library,
            )
            candidate = result["candidate"]
            self.assertFalse(candidate["blockers"])
            confirm_candidate(
                candidate["candidate_id"],
                candidate["candidate_sha256"],
                "确认解析结果",
                library_dir=library,
            )
            save_candidate(candidate["candidate_id"], "范围母版", library_dir=library)
            source_doc = Document()
            source_doc.add_paragraph("新标题")
            source_doc.add_paragraph("新正文")
            source_doc.save(source)
            template = get_template("范围母版", library)
            format_with_template(source, output, template, overwrite=True)
            texts = [paragraph.text for paragraph in Document(output).paragraphs]
            self.assertIn("固定封面", texts)
            self.assertIn("新标题", texts)
            self.assertIn("新正文", texts)
            self.assertNotIn("示例标题", texts)
            self.assertNotIn("示例正文", texts)

    def test_unsupported_doc_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            source = temp / "legacy.doc"
            source.write_bytes(b"legacy")
            with self.assertRaises(FormatterError):
                inspect_template(
                    source, "style", library_dir=temp / "library"
                )

    def test_nested_and_merged_tables_remain_structurally_identical(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            source = temp / "source.docx"
            doc = Document()
            doc.add_paragraph("关于表格的报告")
            outer = doc.add_table(rows=2, cols=2)
            outer.cell(0, 0).text = "外"
            outer.cell(0, 1).text = "表"
            outer.cell(1, 0).merge(outer.cell(1, 1))
            nested = outer.cell(0, 0).add_table(rows=1, cols=1)
            nested.cell(0, 0).text = "嵌套"
            doc.save(source)
            before = document_manifest(Document(source))
            template = get_template("gb9704", temp / "unused")
            formatted = Document(source)
            result = format_style_document(formatted, template["profile"])
            after = document_manifest(formatted)
            self.assertEqual(before, after)
            self.assertGreaterEqual(result["nested_table_count"], 1)

    def test_template_cli_reports_json_and_enforces_selection(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            library = Path(temp_dir) / "library"
            command = [
                sys.executable,
                str(ROOT / "scripts" / "template_cli.py"),
                "list",
                "--library-dir",
                str(library),
            ]
            completed = subprocess.run(
                command, check=False, capture_output=True, text=True
            )
            self.assertEqual(completed.returncode, 0)
            payload = json.loads(completed.stdout)
            self.assertEqual(payload["builtin"][0]["id"], "gb9704")
            self.assertTrue(payload["requires_selection"])

    def test_template_cli_full_lifecycle(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            library = temp / "library"
            source = temp / "template.docx"
            make_template(source)
            base = [
                sys.executable,
                str(ROOT / "scripts" / "template_cli.py"),
            ]
            inspected = subprocess.run(
                base
                + [
                    "inspect",
                    str(source),
                    "--mode",
                    "style",
                    "--library-dir",
                    str(library),
                ],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(inspected.returncode, 0, inspected.stderr)
            payload = json.loads(inspected.stdout)
            candidate = payload["candidate"]
            early_save = subprocess.run(
                base
                + [
                    "save",
                    candidate["candidate_id"],
                    "--name",
                    "CLI模板",
                    "--library-dir",
                    str(library),
                ],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(early_save.returncode, 2)
            confirmed = subprocess.run(
                base
                + [
                    "confirm",
                    candidate["candidate_id"],
                    "--candidate-sha256",
                    candidate["candidate_sha256"],
                    "--user-confirmation-text",
                    "确认解析结果",
                    "--library-dir",
                    str(library),
                ],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(confirmed.returncode, 0, confirmed.stderr)
            saved = subprocess.run(
                base
                + [
                    "save",
                    candidate["candidate_id"],
                    "--name",
                    "CLI模板",
                    "--library-dir",
                    str(library),
                ],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(saved.returncode, 0, saved.stderr)
            self.assertEqual(json.loads(saved.stdout)["status"], "saved")

    def test_format_and_verify_cli_end_to_end(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            source = temp / "source.docx"
            output = temp / "output.docx"
            make_source(source)
            format_report = temp / "format.json"
            verify_report = temp / "verify.json"
            format_command = [
                sys.executable,
                str(ROOT / "scripts" / "format_docx.py"),
                str(source),
                str(output),
                "--template",
                "gb9704",
                "--library-dir",
                str(temp / "library"),
                "--report",
                str(format_report),
                "--verify-report",
                str(verify_report),
            ]
            formatted = subprocess.run(
                format_command, check=False, capture_output=True, text=True
            )
            self.assertEqual(formatted.returncode, 0, formatted.stderr)
            self.assertEqual(json.loads(format_report.read_text())["status"], "success")
            self.assertEqual(json.loads(verify_report.read_text())["status"], "passed")
            verify_command = [
                sys.executable,
                str(ROOT / "scripts" / "verify_docx.py"),
                str(output),
                "--reference",
                str(source),
                "--template",
                "gb9704",
                "--library-dir",
                str(temp / "library"),
                "--report",
                str(temp / "verify-again.json"),
            ]
            verified = subprocess.run(
                verify_command, check=False, capture_output=True, text=True
            )
            self.assertEqual(verified.returncode, 0, verified.stderr)
            self.assertEqual(
                json.loads((temp / "verify-again.json").read_text())["status"],
                "passed",
            )


if __name__ == "__main__":
    unittest.main()
