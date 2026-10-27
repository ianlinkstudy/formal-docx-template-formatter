from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import tempfile
from collections.abc import Iterable, Iterator
from contextlib import contextmanager
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

try:
    import fcntl
except ImportError:  # pragma: no cover - Windows fallback
    fcntl = None


PACKAGE_ROOT = Path(__file__).resolve().parents[2]
GB_PROFILE_PATH = PACKAGE_ROOT / "profiles" / "gb9704.json"
DEFAULT_LIBRARY_DIR = Path.home() / ".formal-docx-template-formatter"
MAX_CUSTOM_TEMPLATES = 3

ROLE_NAMES = (
    "title",
    "recipient",
    "metadata",
    "heading1",
    "heading2",
    "heading3",
    "heading4",
    "body",
    "body_item",
    "table_caption",
    "figure_caption",
    "signature",
    "date",
    "formula_display",
)

LEVEL1_RE = re.compile(r"^\s*[一二三四五六七八九十百零〇两]+、")
LEVEL2_RE = re.compile(r"^\s*[（(][一二三四五六七八九十百零〇两]+[）)]")
LEVEL3_RE = re.compile(r"^\s*\d+(?:[.．、])\s*")
LEVEL4_RE = re.compile(r"^\s*[（(]\d+[）)]")
DATE_RE = re.compile(r"^\s*\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日\s*$")
CAPTION_RE = re.compile(r"^\s*(表|图)\s*[A-Za-z0-9一二三四五六七八九十]+")
METADATA_KEYS = (
    "会议时间",
    "会议地点",
    "会议主题",
    "主持人",
    "出席人员",
    "列席人员",
    "缺席人员",
    "记录人",
)
CLOSING_PREFIXES = ("特此", "请审示", "请批示", "妥否", "以上", "请研处", "请函复")
ORGANIZATION_SUFFIXES = (
    "公司",
    "集团",
    "办公室",
    "委员会",
    "中心",
    "部门",
    "研究院",
    "事业部",
    "学校",
    "医院",
    "协会",
)
ANCHOR_ALIASES = {
    "TITLE": "TITLE",
    "标题": "TITLE",
    "RECIPIENT": "RECIPIENT",
    "主送机关": "RECIPIENT",
    "BODY": "BODY",
    "正文": "BODY",
    "SIGNATURE": "SIGNATURE",
    "落款": "SIGNATURE",
    "DATE": "DATE",
    "日期": "DATE",
}
ANCHOR_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")

MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
PIC_NS = "http://schemas.openxmlformats.org/drawingml/2006/picture"
SVG_NS = "http://schemas.microsoft.com/office/drawing/2016/svg/main"


class FormatterError(RuntimeError):
    """User-actionable formatting or storage error."""


class StateError(FormatterError):
    """The template lifecycle state does not allow the requested action."""


def utc_now() -> str:
    return (
        datetime.now(timezone.utc)  # noqa: UP017 - Python 3.10 compatibility
        .replace(microsecond=0)
        .isoformat()
    )


def load_json(path: Path) -> dict[str, Any]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise FormatterError(f"无法读取 JSON：{path}") from exc
    if not isinstance(payload, dict):
        raise FormatterError(f"JSON 顶层必须是对象：{path}")
    return payload


def atomic_write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
    )
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp_name, path)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)


def atomic_copy(source: Path, target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(
        prefix=f".{target.name}.", suffix=target.suffix, dir=target.parent
    )
    os.close(fd)
    try:
        shutil.copyfile(source, temp_name)
        os.replace(temp_name, target)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)


def deep_merge(base: dict[str, Any], override: dict[str, Any]) -> dict[str, Any]:
    result = deepcopy(base)
    for key, value in override.items():
        if isinstance(value, dict) and isinstance(result.get(key), dict):
            result[key] = deep_merge(result[key], value)
        else:
            result[key] = deepcopy(value)
    return result


def stable_hash(payload: Any) -> str:
    encoded = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


@contextmanager
def library_lock(library_dir: Path) -> Iterator[None]:
    library_dir.mkdir(parents=True, exist_ok=True)
    lock_path = library_dir / ".registry.lock"
    handle = lock_path.open("a+")
    try:
        if fcntl is not None:
            fcntl.flock(handle.fileno(), fcntl.LOCK_EX)
        yield
    finally:
        if fcntl is not None:
            fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
        handle.close()


def _content_type(xml: bytes) -> str:
    match = re.search(
        rb'PartName="/word/document.xml"[^>]*ContentType="([^"]+)"', xml
    )
    if not match:
        match = re.search(
            rb'ContentType="([^"]+)"[^>]*PartName="/word/document.xml"', xml
        )
    return match.group(1).decode("ascii") if match else ""


def prepare_docx_copy(source: Path, target: Path) -> Path:
    """Copy DOCX/DOTX into a DOCX-readable working file without touching source."""
    if not source.is_file():
        raise FormatterError(f"文件不存在：{source}")
    suffix = source.suffix.lower()
    if suffix not in {".docx", ".dotx"}:
        raise FormatterError("仅支持 .docx 和 .dotx；请先将 .doc/.docm/PDF 转为 DOCX。")
    try:
        with ZipFile(source) as package:
            content_types = package.read("[Content_Types].xml")
            package.read("word/document.xml")
    except (BadZipFile, KeyError) as exc:
        raise FormatterError("文件不是有效的 DOCX/DOTX，或已损坏、加密。") from exc

    is_template = (
        suffix == ".dotx"
        or "wordprocessingml.template.main+xml" in _content_type(content_types)
    )
    target.parent.mkdir(parents=True, exist_ok=True)
    if not is_template:
        atomic_copy(source, target)
        return target

    fd, temp_name = tempfile.mkstemp(
        prefix=f".{target.stem}.", suffix=".docx", dir=target.parent
    )
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        with ZipFile(source) as zin, ZipFile(
            temp_path, "w", ZIP_DEFLATED
        ) as zout:
            for item in zin.infolist():
                payload = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    payload = payload.replace(
                        b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
                        b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                    )
                zout.writestr(item, payload)
        os.replace(temp_path, target)
    finally:
        if temp_path.exists():
            temp_path.unlink()
    return target


def open_document(path: Path) -> DocumentObject:
    try:
        return Document(str(path))
    except Exception as exc:
        raise FormatterError(f"无法打开 DOCX：{path}") from exc


def iter_block_items(doc: DocumentObject) -> Iterable[Paragraph | Table]:
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def table_signature(table: Table) -> dict[str, Any]:
    grid = [
        int(column.get(qn("w:w")))
        for column in table._tbl.tblGrid.gridCol_lst
        if column.get(qn("w:w"))
    ]
    rows: list[dict[str, Any]] = []
    for row in table._tbl.tr_lst:
        cells: list[dict[str, Any]] = []
        for tc in row.tc_lst:
            tc_pr = tc.tcPr
            span = tc_pr.find(qn("w:gridSpan")) if tc_pr is not None else None
            merge = tc_pr.find(qn("w:vMerge")) if tc_pr is not None else None
            texts = [
                "".join(node.text or "" for node in p.findall(".//" + qn("w:t")))
                for p in tc.findall("./" + qn("w:p"))
            ]
            nested = [
                table_signature(Table(tbl, table._parent))
                for tbl in tc.findall("./" + qn("w:tbl"))
            ]
            cells.append(
                {
                    "texts": texts,
                    "grid_span": int(span.get(qn("w:val"))) if span is not None else 1,
                    "vertical_merge": (
                        merge.get(qn("w:val"), "continue")
                        if merge is not None
                        else None
                    ),
                    "nested_tables": nested,
                }
            )
        rows.append({"cells": cells})
    return {
        "row_count": len(rows),
        "column_count": len(grid),
        "rows": rows,
    }


def _header_footer_signature(doc: DocumentObject) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for index, section in enumerate(doc.sections):
        for kind, container in (
            ("header", section.header),
            ("first_header", section.first_page_header),
            ("even_header", section.even_page_header),
            ("footer", section.footer),
            ("first_footer", section.first_page_footer),
            ("even_footer", section.even_page_footer),
        ):
            result.append(
                {
                    "section": index,
                    "kind": kind,
                    "paragraphs": [p.text for p in container.paragraphs],
                    "tables": [table_signature(t) for t in container.tables],
                }
            )
    return result


def _canonical_xml_hashes(xml: str, tag: str) -> list[str]:
    from lxml import etree

    root = etree.fromstring(xml.encode("utf-8"))
    hashes = []
    for node in root.iter():
        if etree.QName(node).localname == tag:
            payload = etree.tostring(node, method="c14n")
            hashes.append(hashlib.sha256(payload).hexdigest())
    return sorted(hashes)


def _package_part_hashes(doc: DocumentObject) -> dict[str, list[str]]:
    images: list[str] = []
    embeddings: list[str] = []
    package = doc.part.package
    seen: set[str] = set()
    for part in package.iter_parts():
        partname = str(part.partname)
        if partname in seen:
            continue
        seen.add(partname)
        blob = part.blob
        digest = hashlib.sha256(blob).hexdigest()
        if partname.startswith("/word/media/"):
            images.append(digest)
        elif partname.startswith("/word/embeddings/"):
            embeddings.append(digest)
    return {"images": sorted(images), "embeddings": sorted(embeddings)}


def document_manifest(doc: DocumentObject) -> dict[str, Any]:
    blocks: list[dict[str, Any]] = []
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            blocks.append({"kind": "paragraph", "text": block.text})
        else:
            blocks.append({"kind": "table", "signature": table_signature(block)})

    xml_parts = [doc.element.xml]
    for section in doc.sections:
        for container in (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            xml_parts.append(container._element.xml)
    all_xml = "\n".join(xml_parts)
    formula_hashes: list[str] = []
    object_hashes: list[str] = []
    for xml in xml_parts:
        formula_hashes.extend(_canonical_xml_hashes(xml, "oMath"))
        object_hashes.extend(_canonical_xml_hashes(xml, "object"))
    floating_count = (
        all_xml.count("<wp:anchor")
        + all_xml.count("<w:drawing")
        - all_xml.count("<wp:inline")
    )
    return {
        "blocks": blocks,
        "section_count": len(doc.sections),
        "headers_footers": _header_footer_signature(doc),
        "formula_hashes": sorted(formula_hashes),
        "object_hashes": sorted(object_hashes),
        "part_hashes": _package_part_hashes(doc),
        "inline_image_count": len(doc.inline_shapes),
        "floating_drawing_count": max(0, floating_count),
    }


def manifest_hash(doc: DocumentObject) -> str:
    return stable_hash(document_manifest(doc))


def is_metadata_paragraph(text: str) -> bool:
    stripped = text.strip().replace(" ", "")
    return any(
        stripped.startswith((key + "：", key + ":")) for key in METADATA_KEYS
    )


def _is_closing(text: str) -> bool:
    return any(text.strip().startswith(prefix) for prefix in CLOSING_PREFIXES)


def _looks_like_organization(text: str) -> bool:
    stripped = text.strip()
    if not stripped or len(stripped) > 80 or _is_closing(stripped):
        return False
    if any(stripped.endswith(suffix) for suffix in ORGANIZATION_SUFFIXES):
        return True
    return not re.search(r"[。；;！!？?]", stripped) and len(stripped) <= 30


def _numbered_body_item(text: str, role: str) -> bool:
    stripped = text.strip()
    content = stripped
    for pattern in (LEVEL2_RE, LEVEL3_RE, LEVEL4_RE):
        match = pattern.match(stripped)
        if match:
            content = stripped[match.end() :].strip()
            break
    if role in {"heading3", "heading4"}:
        return any(mark in content for mark in "。；;！!？?")
    if role == "heading2":
        return (
            "。" in content
            and content.count("。") == 1
            and content.endswith("。")
            and len(content) > 20
        )
    return False


def _outline_role(paragraph: Paragraph) -> str | None:
    p_pr = paragraph._p.pPr
    if p_pr is not None:
        outline = p_pr.find(qn("w:outlineLvl"))
        if outline is not None and outline.get(qn("w:val")) in {"0", "1", "2", "3"}:
            return f"heading{int(outline.get(qn('w:val'))) + 1}"
    style_name = paragraph.style.name if paragraph.style is not None else ""
    match = re.search(r"(?:Heading|标题)\s*([1-4])", style_name, re.IGNORECASE)
    if match:
        return f"heading{match.group(1)}"
    return None


def _contains_display_formula(paragraph: Paragraph) -> bool:
    xml = paragraph._p.xml
    if "<m:oMath" not in xml:
        return False
    text_without_math = "".join(
        node.text or "" for node in paragraph._p.findall(".//" + qn("w:t"))
    )
    text_without_math = text_without_math.strip()
    return not text_without_math or bool(re.fullmatch(r"[（(]?\d+[）)]?", text_without_math))


def classify_roles(doc: DocumentObject) -> list[str]:
    paragraphs = doc.paragraphs
    roles = ["body"] * len(paragraphs)
    title_index: int | None = None
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.strip():
            roles[index] = "title"
            title_index = index
            break

    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if is_metadata_paragraph(text):
            roles[index] = "metadata"

    if title_index is not None:
        for index in range(title_index + 1, min(len(paragraphs), title_index + 8)):
            text = paragraphs[index].text.strip()
            if not text or roles[index] == "metadata":
                continue
            if text.endswith(("：", ":")) and len(text) <= 80:
                roles[index] = "recipient"
            break

    for index, paragraph in enumerate(paragraphs):
        if roles[index] in {"title", "recipient", "metadata"}:
            continue
        text = paragraph.text.strip()
        role = _outline_role(paragraph)
        if LEVEL1_RE.match(text):
            role = "heading1"
        elif LEVEL2_RE.match(text):
            role = "heading2"
        elif LEVEL3_RE.match(text):
            role = "heading3"
        elif LEVEL4_RE.match(text):
            role = "heading4"
        if role in {"heading2", "heading3", "heading4"} and _numbered_body_item(
            text, role
        ):
            role = "body_item"
        if role is not None:
            roles[index] = role
        if _contains_display_formula(paragraph):
            roles[index] = "formula_display"
        elif CAPTION_RE.match(text):
            roles[index] = "figure_caption" if text.lstrip().startswith("图") else "table_caption"

    paragraph_by_element = {p._p: i for i, p in enumerate(paragraphs)}
    previous: Paragraph | None = None
    for block in iter_block_items(doc):
        if isinstance(block, Table) and previous is not None:
            previous_index = paragraph_by_element.get(previous._p)
            if (
                previous_index is not None
                and re.match(r"^\s*表\s*[A-Za-z0-9一二三四五六七八九十]+", previous.text)
            ):
                roles[previous_index] = "table_caption"
        if isinstance(block, Paragraph):
            previous = block

    date_index = None
    for index in range(len(paragraphs) - 1, max(title_index or 0, len(paragraphs) - 12) - 1, -1):
        if DATE_RE.match(paragraphs[index].text):
            roles[index] = "date"
            date_index = index
            break
    if date_index is not None:
        for index in range(date_index - 1, max(title_index or 0, date_index - 6) - 1, -1):
            if _looks_like_organization(paragraphs[index].text):
                roles[index] = "signature"
                break
    return roles


def role_map(doc: DocumentObject) -> dict[str, list[int]]:
    roles = classify_roles(doc)
    mapping = {role: [] for role in ROLE_NAMES}
    for index, role in enumerate(roles):
        if role in mapping and doc.paragraphs[index].text.strip():
            mapping[role].append(index)
    return mapping


def normalize_anchor(value: str) -> str | None:
    return ANCHOR_ALIASES.get(value.strip().upper()) or ANCHOR_ALIASES.get(value.strip())


def extract_anchor_names(paragraphs: Iterable[Paragraph]) -> list[str]:
    names: list[str] = []
    for paragraph in paragraphs:
        for match in ANCHOR_RE.finditer(paragraph.text):
            names.append(match.group(1))
    return names


def section_values(section: Any) -> dict[str, Any]:
    return {
        "width_mm": round(section.page_width.mm, 2),
        "height_mm": round(section.page_height.mm, 2),
        "orientation": (
            "landscape"
            if section.orientation == WD_ORIENT.LANDSCAPE
            else "portrait"
        ),
        "margin_top_mm": round(section.top_margin.mm, 2),
        "margin_bottom_mm": round(section.bottom_margin.mm, 2),
        "margin_left_mm": round(section.left_margin.mm, 2),
        "margin_right_mm": round(section.right_margin.mm, 2),
    }


def safe_json(value: Any) -> Any:
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, dict):
        return {str(key): safe_json(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [safe_json(item) for item in value]
    return value
