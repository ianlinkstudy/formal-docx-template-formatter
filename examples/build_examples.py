#!/usr/bin/env python3
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
from base64 import b64decode
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
EXAMPLES = ROOT / "examples"
PYTHON = sys.executable

PNG_BYTES = b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "/A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def run(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [PYTHON, *args],
        cwd=ROOT,
        check=True,
        text=True,
        capture_output=True,
    )


def build_source(path: Path) -> None:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "源文档页眉"
    doc.sections[0].footer.paragraphs[0].text = "源文档页脚"
    doc.add_paragraph("关于模板排版示例项目的报告")
    doc.add_paragraph("示例项目组：")
    doc.add_paragraph("一、总体情况")
    doc.add_paragraph(
        "本示例用于展示模板优先、内容保真的 DOCX 排版流程，所有信息均为虚构。"
    )
    doc.add_paragraph("（一）工作进展。项目已完成模板选择与正文核验。")
    doc.add_paragraph("1. 完成情况")

    table = doc.add_table(rows=3, cols=3)
    values = (
        ("事项", "负责人", "状态"),
        ("模板解析", "示例甲", "完成"),
        ("正文核验", "示例乙", "完成"),
    )
    for row_index, row_values in enumerate(values):
        for column_index, value in enumerate(row_values):
            table.cell(row_index, column_index).text = value
    table.cell(2, 1).merge(table.cell(2, 2))

    doc.add_paragraph("二、下一步")
    doc.add_paragraph(
        "由用户明确选择模板后，再生成新文件；原始内容、数字、标点与表格结构保持不变。"
    )
    doc.add_picture(BytesIO(PNG_BYTES))
    doc.inline_shapes[-1].width = Inches(0.35)
    doc.inline_shapes[-1].height = Inches(0.18)
    doc.add_paragraph("特此报告。")
    signature = doc.add_paragraph("鼓翼成长星球示例项目组")
    signature.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    date = doc.add_paragraph("2026年10月27日")
    date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(16)
    doc.save(path)


def build_style_template(path: Path) -> None:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "示例公司页眉"
    doc.sections[0].footer.paragraphs[0].text = "示例公司页脚"
    doc.add_paragraph("公司标准报告标题")
    doc.add_paragraph("一、部门情况")
    heading = doc.add_paragraph("（一）工作进展。示例公司标准正文从这里开始。")
    heading.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    table = doc.add_table(rows=2, cols=3)
    for row_index, values in enumerate(
        (("序号", "事项", "状态"), ("1", "示例事项", "完成"))
    ):
        for column_index, value in enumerate(values):
            table.cell(row_index, column_index).text = value
    table.cell(1, 1).merge(table.cell(1, 2))
    doc.add_paragraph("示例公司名称")
    doc.add_paragraph("2026年10月27日")
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(16)
    doc.save(path)


def build_shell_template(path: Path) -> None:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "示例公司 · 固定页眉"
    doc.sections[0].footer.paragraphs[0].text = "第 1 页 · 示例模板"
    doc.add_paragraph("示例公司标准报告")
    doc.add_paragraph("模板固定封面内容（不会写入正文）")
    doc.add_paragraph("{{TITLE}}")
    doc.add_paragraph("{{RECIPIENT}}")
    doc.add_paragraph("{{BODY}}")
    doc.add_paragraph("{{SIGNATURE}}")
    doc.add_paragraph("{{DATE}}")
    doc.save(path)


def inspect_confirm_save(template: Path, library: Path, name: str) -> None:
    inspected = json.loads(
        run(
            "scripts/template_cli.py",
            "inspect",
            str(template),
            "--mode",
            "auto",
            "--library-dir",
            str(library),
        ).stdout
    )
    candidate = inspected["candidate"]
    if candidate["blockers"]:
        raise RuntimeError(f"示例模板存在阻断项：{candidate['blockers']}")
    run(
        "scripts/template_cli.py",
        "confirm",
        candidate["candidate_id"],
        "--candidate-sha256",
        candidate["candidate_sha256"],
        "--user-confirmation-text",
        "确认解析结果",
        "--library-dir",
        str(library),
    )
    run(
        "scripts/template_cli.py",
        "save",
        candidate["candidate_id"],
        "--name",
        name,
        "--library-dir",
        str(library),
    )


def main() -> int:
    before_dir = EXAMPLES / "before"
    template_dir = EXAMPLES / "templates"
    after_dir = EXAMPLES / "after"
    reports_dir = EXAMPLES / "reports"
    for directory in (before_dir, template_dir, after_dir, reports_dir):
        directory.mkdir(parents=True, exist_ok=True)

    source = before_dir / "01-project-report.docx"
    style_template = template_dir / "01-style-template.docx"
    shell_template = template_dir / "02-shell-template.docx"
    build_source(source)
    build_style_template(style_template)
    build_shell_template(shell_template)

    with tempfile.TemporaryDirectory() as temp_dir:
        library = Path(temp_dir) / "library"
        inspect_confirm_save(style_template, library, "示例公司样式")
        inspect_confirm_save(shell_template, library, "示例公司母版")

        output = after_dir / "01-gb9704-project-report.docx"
        run(
            "scripts/format_docx.py",
            source.relative_to(ROOT).as_posix(),
            output.relative_to(ROOT).as_posix(),
            "--template",
            "GB/T 9704-2012 参考版",
            "--library-dir",
            str(library),
            "--overwrite",
            "--report",
            (reports_dir / "01-gb9704.format-report.json")
            .relative_to(ROOT)
            .as_posix(),
            "--verify-report",
            (reports_dir / "01-gb9704.verify-report.json")
            .relative_to(ROOT)
            .as_posix(),
        )

        output = after_dir / "02-company-shell-project-report.docx"
        run(
            "scripts/format_docx.py",
            source.relative_to(ROOT).as_posix(),
            output.relative_to(ROOT).as_posix(),
            "--template",
            "示例公司母版",
            "--library-dir",
            str(library),
            "--overwrite",
            "--report",
            (reports_dir / "02-company-shell.format-report.json")
            .relative_to(ROOT)
            .as_posix(),
            "--verify-report",
            (reports_dir / "02-company-shell.verify-report.json")
            .relative_to(ROOT)
            .as_posix(),
        )

    print("示例已生成并完成核验。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
